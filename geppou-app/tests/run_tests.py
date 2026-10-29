"""ブラウザ版（geppou/）の自動テスト。

    python3 geppou-app/tests/run_tests.py          # すべて実行
    python3 geppou-app/tests/run_tests.py 集計     # 名前に「集計」を含むテストだけ

答えの分かっている見本の月報（fixtures.py、すべて架空）を実際のブラウザ（Chromium）で取り込み、
読み取り・集計・未提出・異動・AI用テキスト・PDF・バックアップの結果が答えと一致するかを確かめる。
あわせて、アプリが外部（インターネット）へ一切通信していないことも確かめる。
修正のたびに実行し、すべて ✓ になってから公開する。

必要なもの: Python の playwright（Chromium 入り）, pypdf, reportlab, openpyxl, python-docx, Pillow, IPAゴシック
"""
import base64
import functools
import http.server
import io
import os
import re
import sys
import tempfile
import threading
import time
import traceback
import unicodedata
from urllib.parse import urlparse

from playwright.sync_api import sync_playwright

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
sys.path.insert(0, HERE)
import fixtures  # noqa: E402

TESTS = []


def test(name):
    def deco(fn):
        TESTS.append((name, fn))
        return fn
    return deco


def eq(actual, expected, what=""):
    if actual != expected:
        raise AssertionError(f"{what}: {expected!r} のはずが {actual!r}")


def nfkc(pairs):
    # アプリは全角の記号・英数字を半角にそろえて読む（［かぜ］→[かぜ]）
    return [(unicodedata.normalize("NFKC", n), c) for n, c in pairs]


def ok(cond, what):
    if not cond:
        raise AssertionError(what)


# ----------------------------------------------------------------------- 準備

def start_server():
    handler = functools.partial(QuietHandler, directory=ROOT)
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return srv


class QuietHandler(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a):
        pass


class App:
    """1つのブラウザ（保存領域つき）でアプリを開く。外部への通信はすべて止めて記録する"""

    def __init__(self, browser, base, query=""):
        self.base = base
        self.external = []
        self.errors = []
        self.ctx = browser.new_context(viewport={"width": 1300, "height": 900}, service_workers="block", accept_downloads=True)
        self.ctx.route("**/*", self._route)
        self.page = self.ctx.new_page()
        self.page.on("pageerror", lambda e: self.errors.append(str(e) + (" @ " + (e.stack or "").split("\n")[1].strip() if getattr(e, "stack", None) and len((e.stack or "").split("\n")) > 1 else "")))
        self.open(query)

    def _route(self, route):
        u = urlparse(route.request.url)
        if u.scheme in ("http", "https") and u.netloc != urlparse(self.base).netloc:
            self.external.append(route.request.url)
            route.abort()
        else:
            route.continue_()

    def open(self, query=""):
        self.page.goto(f"{self.base}/geppou/{query}")
        self.page.wait_for_function("typeof state !== 'undefined' && !!document.querySelector('#tabs button') && (window.appReady === true || /pdfjob|tourtab/.test(location.search))")
        if self.page.evaluate("document.querySelector('#notice') && document.querySelector('#notice').open"):
            self.page.check("#noticeOk")
            self.page.click("#noticeBtn")
        self.page.wait_for_timeout(300)

    def js(self, expr, arg=None):
        return self.page.evaluate(expr, arg)

    def tab(self, name):
        self.page.click(f'#tabs button[data-page="{name}"]')
        self.page.wait_for_timeout(400)

    def close(self):
        self.ctx.close()


# ----------------------------------------------------------------------- 個々の読み取り（関数単位）

@test("公館名の正規化（表記ゆれ）")
def t_canonical(app):
    for raw, want in [("在ケニア日本国大使館", "ケニア"), ("在ケニア大使館", "ケニア"), ("ケニア", "ケニア"),
                      ("在フィリピン日本国大使館", "フィリピン"), ("在タイ大使館", "タイ")]:
        eq(app.js("(s) => canonical(s)", raw), want, raw)
    eq(app.js("canonical('在存在しない国大使館')"), None, "存在しない公館")


@test("ファイル名・本文から対象月を読む")
def t_month(app):
    for name, want in [("ケニア_202607_月報.pdf", "2026-07"), ("2026年8月_ウガンダ.docx", "2026-08"),
                       ("令和8年7月 ケニア.docx", "2026-07"), ("R8.7 ケニア.pdf", "2026-07"),
                       ("2026-08/ケニア.pdf", "2026-08"), ("ケニア.pdf", None), ("ケニア_202613.pdf", None)]:
        eq(app.js("(s) => monthFromName(s)", name), want, name)
    eq(app.js("detectMonth('受診者数等調査票\\n（2026年7月分）', 'x.pdf')"), "2026-07", "本文の年月")
    eq(app.js("detectMonth('令和8年9月', 'x.pdf')"), "2026-09", "本文の令和")
    eq(app.js("['月報_2025011.pdf', 'ケニア2025012.pdf', 'x/2025001/a.pdf'].map(monthFromName)"), ["2025-11", "2025-12", "2025-01"], "年4桁＋月3桁のファイル名")
    eq(app.js("['202511', '2025011', '2025/11', '2025-1', '2025年11月', 'R7.11', '令和7年11月', '202511年', '', '202513', 'abc'].map(parseMonthInput)"),
       ["2025-11", "2025-11", "2025-11", "2025-01", "2025-11", "2025-11", "2025-11", "2025-11", "", None, None], "手入力の年月")


@test("甲表（受診者数）の読み取り")
def t_kou(app):
    r = app.js("(t) => parseKou(t)", fixtures.pages_of("ケニア_202607_月報.pdf")[0])
    eq(r["embassy"], "在ケニア日本国大使館", "公館")
    eq(r["officer"], "架空 一郎", "医務官")
    eq(r["counts"], [10, 5, 3, 18], "受診者数")


@test("乙表（病名と件数）の読み取りと疾病分類")
def t_otsu(app):
    r = app.js("(t) => parseOtsu([t])", fixtures.pages_of("ケニア_202607_月報.pdf")[1])
    eq([(d["name"], d["count"]) for d in r], nfkc(fixtures.REPORTS["ケニア_202607_月報.pdf"][4]), "病名と件数")
    cats = app.js("['デング熱','胃腸炎','健康診断','湿疹'].map(classify)")
    ok(all(cats), f"分類できない病名がある: {cats}")
    eq(len(set(cats)), 4, "4つの病名が別々の分類になること")


@test("病名の分類：辞書を広げ、乙表の分類の欄も手がかりにして「分類不能」を減らす")
def t_classify(app):
    want = {"急性気管支炎": "呼吸器", "花粉症": "呼吸器", "逆流性食道炎": "消化器", "虫歯": "消化器", "細菌性腸炎": "感染症", "とびひ": "感染症",
            "溶連菌感染症": "感染症", "白癬": "感染症", "パニック障害": "精神・行動", "抑うつ状態": "精神・行動", "あせも": "皮膚", "蜂窩織炎": "皮膚",
            "ドライアイ": "眼", "ものもらい": "眼", "耳垢栓塞": "耳", "肩関節周囲炎": "筋骨格", "腰椎椎間板ヘルニア": "筋骨格", "足関節捻挫": "外傷・中毒",
            "切創": "外傷・中毒", "時差ぼけ": "外傷・中毒", "月経困難症": "腎尿路生殖器", "前立腺肥大症": "腎尿路生殖器", "片頭痛": "神経",
            "顔面神経麻痺": "神経", "脂質異常症": "内分泌・代謝", "脱水症": "内分泌・代謝", "鉄欠乏性貧血": "血液", "本態性高血圧症": "循環器",
            "動悸": "循環器", "つわり": "妊娠・出産", "大腸ポリープ": "新生物", "発熱": "症状・徴候", "健康相談": "検診・予防接種等(Z)",
            "黄熱ワクチン接種": "検診・予防接種等(Z)", "蕁麻疹": "皮膚"}
    got = app.js("(xs) => Object.fromEntries(xs.map(x => [x, shortCat(classify(x))]))", list(want))
    eq({k: v for k, v in got.items() if v != want[k]}, {}, "辞書で分類できない病名")
    # 辞書にない病名でも、乙表でその病名が書かれていた分類の欄から決める（病名の前の個別のICDコードを優先）
    otsu = ("疾病分類件数表\n在ケニア日本国大使館\n（2026年7月分）\n"
            "1 Ｊ００－Ｊ９９ 呼吸器系の疾患 5 ほげほげ症(3) ふがふが(2)\n"
            "2 Ｍ００－Ｍ９９ 筋骨格系及び結合組織の疾患 4 なぞの痛み(4)\n"
            "3 Ｌ００－Ｌ９９ 皮膚及び皮下組織の疾患 2 Ｈ１０ よくわからない病(2)\n報告者 医務官")
    r = app.js("(t) => parseOtsu([t]).map(d => [d.name, d.count, shortCat(classify(d.name, d.cat))])", otsu)
    eq(r, [["ほげほげ症", 3, "呼吸器"], ["ふがふが", 2, "呼吸器"], ["なぞの痛み", 4, "筋骨格"], ["よくわからない病", 2, "眼"]], "分類の欄からの分類")
    eq(app.js("shortCat(classify('ほげほげ症'))"), "分類不能", "手がかりのないときは分類不能のまま")


@test("丙表から感染症の流行・終息を読む")
def t_hei(app):
    r = app.js("heiDiseases('雨季入りに伴いデング熱が急増し、病床がひっ迫している。')")
    eq([(d["name"], d["trend"]) for d in r], [("デング熱", "流行・増加")], "急増")
    r = app.js("heiDiseases('コレラの新規報告は減少傾向にある。')")
    eq([(d["name"], d["trend"]) for d in r], [("コレラ", "減少・終息")], "減少")
    r = app.js("heiDiseases('鳥インフルエンザの症例が報告された。')")
    eq([d["name"] for d in r], ["鳥インフルエンザ"], "鳥インフルエンザをインフルエンザと二重に数えない")
    r = app.js("heiDiseases('館員が蕁麻疹で受診した。')")
    eq(r, [], "蕁麻疹は感染症（麻疹）にしない")


@test("丙表から巡回健診・緊急移送を読む")
def t_junkai_cases(app):
    j = app.js("detectJunkai('8月18日～21日 南スーダンへ巡回健診を実施し、館員12名を診察した。', 'ウガンダ')")
    eq([(x["from"], x["dest"], x["period"]) for x in j], [("ウガンダ", "南スーダン", "8/18〜8/21")], "巡回健診")
    c = app.js("detectCases('邦人旅行者1名が交通外傷で入院し、南アフリカの病院へ緊急移送した。手術を受け回復した。', 'ケニア')")
    evac = [x for x in c if x["type"] == "evac"]
    eq(len(evac), 1, "緊急移送の件数")
    eq((evac[0]["who"], evac[0]["dest"], evac[0]["outcome"]), ("邦人旅行者", "南アフリカ共和国", "回復・退院"), "移送の内容（転帰は次の文から）")
    eq(app.js("detectCases('邦人の移送は不要と判断した。', 'ケニア').filter(x => x.type === 'evac').length"), 0, "「移送は不要」は移送にしない")


# ----------------------------------------------------------------------- 取り込み → 集計（画面を通して）

def import_fixtures(app, folder):
    app.tab("import")
    app.page.check('input[name="mmode"][value="auto"]')
    files = [os.path.join(folder, f) for f in sorted(os.listdir(folder))]
    app.page.set_input_files("#files", files)
    app.page.wait_for_selector("[data-rsave]")
    app.page.click('[data-rsave="0"]')
    app.page.wait_for_function("Object.keys(state.rosters).length === 1")
    app.page.click("#start")
    app.page.wait_for_selector("#saveDrafts", timeout=60000)
    drafts = app.js("drafts.map(d => ({f: d.filename, w: d.warnings, post: d.post, month: d.month, format: d.format}))")
    app.page.click("#saveDrafts")
    app.page.wait_for_function("Object.keys(state.reports).length >= 6")
    return drafts


@test("取り込み：読み取り中に他のタブを開いても止まらない／要確認以外をすべて保存／同じ月報は上書きかスキップを確認")
def t_import_bg(browser, base):
    app = App(browser, base)
    try:
        folder = fixtures.build(tempfile.mkdtemp(prefix="geppou_bg_"))
        app.tab("import")
        app.page.check('input[name="mmode"][value="auto"]')
        app.page.set_input_files("#files", [os.path.join(folder, f) for f in sorted(os.listdir(folder)) if not f.endswith(".xlsx")])
        app.js("window._toasts = []; const _t = toast; toast = (m, ...a) => { _toasts.push(m); return _t(m, ...a); }")
        app.page.click("#start")
        app.page.wait_for_function("importBusy")
        ok(app.js("document.querySelector('#tabs button[data-page=import]').classList.contains('busy')"), "読み取り中の印がタブに出ない")
        for name in ["dashboard", "stats", "map", "summary"]:
            app.tab(name)
        app.page.wait_for_function("!importBusy", timeout=60000)
        eq(app.js("drafts.length"), 6, "他のタブを開いた後の読み取り件数")
        ok(app.js("drafts.filter(d => /受診者数/.test(d.kou || '')).length") >= 5, "確認の詳細に甲表の文字がない")
        ok(app.js("!document.querySelector('#tabs button[data-page=import]').classList.contains('busy')"), "読み取り後も印が消えない")
        ok(any("読み取りが終わりました" in (m or "") for m in app.js("_toasts.map(String)")), "終わったお知らせが出ない")
        app.tab("import")
        ok(app.page.locator("#saveDrafts").count() == 1, "取り込みタブに戻っても確認の表が出ない")
        ok("6件を読み取りました" in app.js("document.querySelector('#progress').textContent"), "取り込みタブに結果の件数が出ない")
        # 対象月の欄に「202511」と入れると「2025年11月」になり、2025-11 として扱う（あとで元に戻す）
        f = app.page.locator('#drafts tr[data-i="0"] input[data-f="month"]')
        orig = f.input_value()
        f.fill("202511"); f.press("Tab")
        eq([f.input_value(), app.js("(syncDrafts(), drafts[0].month)")], ["2025年11月", "2025-11"], "対象月の手入力")
        f.fill(orig); f.press("Tab"); app.js("syncDrafts()")
        # 人数の欄：全角の数字で入れても半角に直り、計も自動で計算する
        row = '#drafts tr[data-i="0"] '
        keep = [app.page.input_value(row + f'input[data-f="{k}"]') for k in ("staff", "family", "other", "total")]
        app.page.fill(row + 'input[data-f="staff"]', "１２")
        app.page.fill(row + 'input[data-f="family"]', "３")
        app.page.fill(row + 'input[data-f="other"]', "0")
        eq([app.page.input_value(row + f'input[data-f="{k}"]') for k in ("staff", "family", "total")], ["12", "3", "15"], "全角の数字を半角に直す")
        for k, v in zip(("staff", "family", "other", "total"), keep): app.page.fill(row + f'input[data-f="{k}"]', v)
        app.js("syncDrafts()")
        # 「要確認以外をすべて保存」：警告のない5件だけ保存し、要確認のタイは表に残る
        ok("5件" in app.page.inner_text("#saveOkDrafts"), "ボタンの件数")
        app.page.click("#saveOkDrafts")
        app.page.wait_for_function("Object.keys(state.reports).length === 5 && drafts.length === 1")
        eq(app.js("drafts.map(d => d.filename)"), ["タイ_202607_月報.pdf"], "表に残った要確認の行")
        ok(app.page.locator("#saveDrafts").count() == 1 and app.page.locator("#saveOkDrafts").count() == 0, "残った行の確認の表")
        # 再読み込みしても、保存していない読み取り結果（直した内容も）は残る
        app.page.fill('#drafts tr[data-i="0"] input[data-f="officer"]', "直した 名前")
        app.page.set_input_files("#files", [os.path.join(folder, "ウガンダ_202607_月報.pdf")])
        app.page.wait_for_timeout(600)
        app.open()
        app.page.wait_for_selector("#saveDrafts")
        eq(app.js("[drafts.length, drafts[0].filename, drafts[0].officer, page]"), [1, "タイ_202607_月報.pdf", "直した 名前", "import"], "再読み込み後の確認の表")
        eq(app.js("pending.map(p => [p.file.name, p.file.size > 0])"), [["ウガンダ_202607_月報.pdf", True]], "再読み込み後の読み取る前のファイル")
        # 新しく入れたファイルは一覧の一番上に
        app.page.set_input_files("#files", [os.path.join(folder, "ケニア_202608_月報.pdf")])
        eq(app.js("pending.map(p => p.file.name)"), ["ケニア_202608_月報.pdf", "ウガンダ_202607_月報.pdf"], "新しいファイルが一番上にない")
        ok(app.js("pending[0].file.arrayBuffer().then(b => b.byteLength)") > 0, "戻したファイルを読めない")
        app.js("pending = []; persistDrafts()")
        app.js("window._toasts = []; const _t2 = toast; toast = (m, ...a) => { _toasts.push(m); return _t2(m, ...a); }")
        app.page.click("#saveDrafts")
        app.page.wait_for_function("Object.keys(state.reports).length === 6 && drafts.length === 0")
        # 同じファイルをもう一度：一覧には重ねて入らず、保存のときに上書き／スキップを聞く
        files = [os.path.join(folder, f) for f in sorted(os.listdir(folder)) if f.startswith("ケニア")]
        app.page.set_input_files("#files", files)
        app.page.set_input_files("#files", files)
        ok(any("重ねて追加しませんでした" in (m or "") for m in app.js("_toasts.map(String)")), "同じファイルを重ねて追加したお知らせがない")
        eq(app.js("pending.length"), 2, "一覧のファイル数")
        app.js("state.reports['2026-07|ケニア'].hei = '保存済みの内容'")
        app.page.click("#start")
        app.page.wait_for_selector("#saveDrafts", timeout=60000)
        app.page.click("#saveDrafts")
        app.page.wait_for_selector("dialog.dupdlg[open]")
        ok("2件" in app.page.inner_text("dialog.dupdlg"), "確認の件数")
        app.page.click('dialog.dupdlg [data-ans="skip"]')
        app.page.wait_for_function("drafts.length === 0")
        eq(app.js("state.reports['2026-07|ケニア'].hei"), "保存済みの内容", "スキップしたのに上書きされた")
        app.page.set_input_files("#files", files[:1])
        app.page.click("#start")
        app.page.wait_for_selector("#saveDrafts", timeout=60000)
        app.page.click("#saveDrafts")
        app.page.click('dialog.dupdlg [data-ans="overwrite"]')
        app.page.wait_for_function("drafts.length === 0")
        ok("コレラ" in app.js("state.reports['2026-07|ケニア'].hei"), "上書きされていない")
        eq(app.js("Object.keys(state.reports).length"), 6, "上書き後の件数")
        eq(app.errors, [], "画面のエラー")
    finally:
        app.close()


@test("取り込み：1件ずつ読み取る・消去／選んだものだけ読み取る・保存")
def t_pick(browser, base):
    app = App(browser, base)
    try:
        folder = fixtures.build(tempfile.mkdtemp(prefix="geppou_pick_"))
        names = ["ケニア_202607_月報.pdf", "ケニア_202608_月報.pdf", "ウガンダ_202607_月報.pdf", "タイ_202607_月報.pdf"]
        app.tab("import")
        app.page.check('input[name="mmode"][value="auto"]')
        app.page.set_input_files("#files", [os.path.join(folder, n) for n in names])
        eq(app.page.locator("[data-pread]").count(), 4, "1件ずつの「読み取る」ボタン")
        # 一覧から1件消去 → 3件
        app.page.click('[data-rm="3"]')
        eq(app.js("pending.map(p => p.file.name)"), names[:3], "消去した後の一覧")
        # 1件のチェックを外す → 「選択した 2件を読み取る」
        app.page.uncheck('[data-psel="2"]')
        eq(app.page.inner_text("#start"), "選択した 2件を読み取る", "読み取りボタンの表示")
        app.page.click("#start")
        app.page.wait_for_selector("#saveDrafts", timeout=60000)
        eq(sorted(app.js("drafts.map(d => d.filename)")), sorted(names[:2]), "選んだ2件だけ読み取る")
        eq(app.js("pending.map(p => p.file.name)"), [names[2]], "選ばなかったファイルは一覧に残る")
        # 残った1件をその行の「読み取る」で → 表は3件に増える
        app.page.click('[data-pread="0"]')
        app.page.wait_for_function("!importBusy && drafts.length === 3", timeout=60000)
        eq(app.js("pending.length"), 0, "読み取った後の一覧")
        # 確認の表：1行消去 → 2件、1行だけ保存 → 1件保存・1件残る、選んで保存
        i = app.js("drafts.findIndex(d => d.filename === 'ケニア_202608_月報.pdf')")
        app.page.click(f'[data-ddel="{i}"]')
        eq(app.js("drafts.length"), 2, "表から消去した後の件数")
        i = app.js("drafts.findIndex(d => d.filename === 'ウガンダ_202607_月報.pdf')")
        app.page.fill(f'tr[data-i="{i}"] [data-f="officer"]', "直した 名前")
        # 公館の欄を押すと、名前が入っていても候補の一覧がすべて出る。文字を入れると絞り込み、押すと選べる
        app.page.click(f'tr[data-i="{i}"] [data-f="post"]')
        ok(app.js("document.querySelectorAll('.combo-list .combo-item').length") > 100, "欄を押しても候補の一覧が出ない")
        app.page.fill(f'tr[data-i="{i}"] [data-f="post"]', "ホンジ")
        eq(app.js("[...document.querySelectorAll('.combo-list .combo-item')].map(x => x.firstChild.textContent)"), ["ホンジュラス"], "入力で絞り込み")
        app.page.click(".combo-list .combo-item")
        eq(app.page.input_value(f'tr[data-i="{i}"] [data-f="post"]'), "ホンジュラス", "候補を押して選ぶ")
        eq(app.js("!!document.querySelector('.combo-list')"), False, "選んだら一覧が閉じる")
        # 公館名は入力して検索できる（英語名・都市名でもよい）。ホンジュラスも選べる
        eq(app.js("[...document.querySelectorAll('#postList option')].some(o => o.value === 'ホンジュラス')"), True, "候補にホンジュラスがない")
        app.page.fill(f'tr[data-i="{i}"] [data-f="post"]', "Honduras")
        app.page.dispatch_event(f'tr[data-i="{i}"] [data-f="post"]', "change")
        eq(app.page.input_value(f'tr[data-i="{i}"] [data-f="post"]'), "ホンジュラス", "英語名から公館名へ")
        app.page.fill(f'tr[data-i="{i}"] [data-f="post"]', "在ウガンダ日本国大使館")
        app.page.dispatch_event(f'tr[data-i="{i}"] [data-f="post"]', "change")
        eq(app.page.input_value(f'tr[data-i="{i}"] [data-f="post"]'), "ウガンダ", "正式名から公館名へ")
        # 職員・家族・その他を手で入れると計は自動で計算される
        app.page.fill(f'tr[data-i="{i}"] [data-f="staff"]', "8")
        app.page.fill(f'tr[data-i="{i}"] [data-f="family"]', "4")
        app.page.fill(f'tr[data-i="{i}"] [data-f="other"]', "")
        eq(app.page.input_value(f'tr[data-i="{i}"] [data-f="total"]'), "12", "計の自動計算")
        app.page.click(f'[data-dsave="{i}"]')
        app.page.wait_for_function("Object.keys(state.reports).length === 1 && drafts.length === 1")
        eq(app.js("state.reports['2026-07|ウガンダ'].officer"), "直した 名前", "表で直した内容で保存")
        eq(app.js("[state.reports['2026-07|ウガンダ'].staff, state.reports['2026-07|ウガンダ'].other, state.reports['2026-07|ウガンダ'].total]"), [8, None, 12], "自動計算した計で保存")
        eq(app.js("drafts.map(d => d.filename)"), ["ケニア_202607_月報.pdf"], "1件保存した後の表")
        app.page.click("#dselAll")
        ok(app.page.is_disabled("#saveSelDrafts"), "何も選んでいないのに保存ボタンが押せる")
        app.page.check(".dsel")
        app.page.click("#saveDrafts")
        app.page.wait_for_function("Object.keys(state.reports).length === 2 && drafts.length === 0")
        eq(app.errors, [], "画面のエラー")
    finally:
        app.close()


@test("配置表：PDFも読める（配置表の欄へのドロップ・月報と一緒に入れても振り分け）")
def t_roster_pdf(browser, base):
    app = App(browser, base)
    try:
        d = tempfile.mkdtemp(prefix="geppou_roster_")
        named, plain = os.path.join(d, "医務官配置表_202608.pdf"), os.path.join(d, "202608_一覧.pdf")
        fixtures.make_roster_pdf(named)
        fixtures.make_roster_pdf(plain)
        app.tab("import")
        want = [["ケニア", "見本 花子", "みほん はなこ", True], ["ウガンダ", "架空 一郎", "かくう いちろう", True], ["フィリピン", "検証 六子", "けんしょう むつこ", True],
                ["タイ", "試作 七美", "しさく ななみ", True], ["スーダン", "", "", False]]
        rows = "rosterDrafts.map(r => [r.month, r.rows.map(x => [x.post, x.officer, x.kana, x.staffed])])"
        # 1) 配置表の欄にドラッグ＆ドロップ
        app.page.evaluate("""async (b64) => {
            const u8 = Uint8Array.from(atob(b64), c => c.charCodeAt(0));
            const dt = new DataTransfer(); dt.items.add(new File([u8], '医務官配置表_202608.pdf', {type: 'application/pdf'}));
            document.querySelector('#rosterDrop').dispatchEvent(new DragEvent('drop', {dataTransfer: dt, bubbles: true, cancelable: true})); }""",
            base64.b64encode(open(named, "rb").read()).decode())
        app.page.wait_for_function("rosterDrafts.length === 1")
        eq(app.js(rows), [["2026-08", want]], "ドロップした PDF の配置表")
        eq(app.js("rosterDrafts[0].monthFrom"), "ファイル名", "適用月の読み取り元")
        app.js("rosterDrafts = []")
        # 2) 月報の欄に入れた PDF：名前に「配置」があれば配置表、なくても中身で振り分ける
        app.page.check('input[name="mmode"][value="auto"]')
        app.page.set_input_files("#files", [named])
        app.page.wait_for_function("rosterDrafts.length === 1")
        eq(app.js("pending.length"), 0, "名前に「配置」がある PDF が月報の一覧に入った")
        app.js("rosterDrafts = []")
        app.page.set_input_files("#files", [plain])
        app.page.click("#start")
        app.page.wait_for_function("!importBusy && rosterDrafts.length === 1", timeout=60000)
        eq(app.js("drafts.length"), 0, "中身が配置表の PDF を月報として読んだ")
        eq(app.js(rows), [["2026-08", want]], "中身で振り分けた配置表")
        app.page.click('[data-rsave="0"]')
        app.page.wait_for_function("Object.keys(state.rosters).length === 1")
        eq(app.errors, [], "画面のエラー")
    finally:
        app.close()


@test("配置表：公館の欄の「国名」は大使館、「◯◯総」は総領事館と読む")
def t_roster_abbrev(app):
    txt = "\n".join(["医務官配置表（2026年8月1日現在）", "No 地域 公館 氏名 ふりがな 備考", "1 アフリカ ケニア 見本 花子 みほん はなこ", "2 アジア 広州総 雛形 八郎",
                     "3 アジア 上海 総 試作 七美", "4 アジア 香港（総） 例示 三郎", "5 アジア 総瀋陽 原案 九子", "6 アジア 中国 草稿 十郎", "7 北米 ニューヨーク総 素案 一美",
                     "8 アフリカ スーダン 退避中", "9 在タイ日本国大使館 試作 七美"])
    r = app.js("(t) => parseRosterText(t).map(r => [r.post, r.officer, r.staffed])", txt)
    eq(r, [["ケニア", "見本 花子", True], ["広州総領事館", "雛形 八郎", True], ["上海総領事館", "試作 七美", True], ["香港総領事館", "例示 三郎", True],
           ["瀋陽総領事館", "原案 九子", True], ["中国", "草稿 十郎", True], ["ニューヨーク総領事館", "素案 一美", True], ["スーダン", "", False], ["タイ", "試作 七美", True]], "公館の読み取り")
    eq(app.js("rosterPostName('広州総')"), "広州総領事館", "Excelの欄でも同じ読み方")


@test("取り込み：6件すべて読めて、警告は合計不一致の1件だけ")
def t_import(app):
    drafts = {d["f"]: d for d in app.drafts}
    eq(len(drafts), 6, "読み取った件数")
    for f, d in drafts.items():
        if f.startswith("タイ"):
            ok(any("合計が一致しません" in w for w in d["w"]), f"{f}: 合計不一致の警告が出ない {d['w']}")
        else:
            eq(d["w"], [], f"{f} の警告")
    eq(drafts["フィリピン_202608_月報.pdf"]["post"], "フィリピン", "パスワード付きPDF（imukanmd）が開ける")
    eq(drafts["2026年8月_ウガンダ.docx"]["month"], "2026-08", "Wordの対象月")
    eq(app.js("Object.keys(state.reports).sort()"),
       ["2026-07|ウガンダ", "2026-07|ケニア", "2026-07|タイ", "2026-08|ウガンダ", "2026-08|ケニア", "2026-08|フィリピン"], "保存された月報")


@test("集計：受診者数・病名の件数が見本と一致")
def t_totals(app):
    for fname, (_, officer, period, counts, dis, _) in fixtures.REPORTS.items():
        y, m = period.replace("月", "").split("年")
        post = app.js("(e) => canonical(e)", fixtures.REPORTS[fname][0])
        r = app.js("(k) => state.reports[k]", f"{y}-{int(m):02d}|{post}")
        eq([r["staff"], r["family"], r["other"], r["total"]], list(counts), f"{fname} の受診者数")
        eq(r["officer"], officer, f"{fname} の医務官")
        eq(sorted((d["name"], d["count"]) for d in r["diseases"]), sorted(nfkc(dis)), f"{fname} の乙表")
    eq(app.js("reportsOf('2026-08').reduce((s, r) => s + r.total, 0)"), 20 + 11 + 15, "8月の受診者数の合計")
    eq(app.js("diseasesOf('2026-08').filter(d => d.name === 'デング熱').reduce((s, d) => s + d.count, 0)"), 5, "8月のデング熱")


@test("未提出公館：配置表と突合してタイだけ")
def t_missing(app):
    eq(app.js("missingPosts('2026-08').map(r => r.post)"), ["タイ"], "8月の未提出")
    miss = app.js("submitMatrix(0).miss.map(x => x.month + ' ' + x.post)")
    ok("2026-08 タイ" in miss, f"提出状況の足りない月報: {miss}")
    # 古い月の配置表があっても、公館の顔ぶれは最新の配置表のものを全部の月に当てはめる
    eq(app.js("(() => { state.rosters['2026-06'] = [{post: 'ケニア', officer: 'x', staffed: true}]; const r = [...expectedPosts('2026-07')].sort(); delete state.rosters['2026-06']; return r; })()"),
       sorted(["ウガンダ", "ケニア", "スーダン", "タイ", "フィリピン"]), "過去の月にも最新の配置表の公館")
    # 配置表がないときは、公館一覧の医務官のいる公館すべてを基準にする（月報を出した公館だけだと「すべて提出」になってしまう）
    nr = app.js("(() => { const keep = state.rosters; state.rosters = {}; const r = submitMatrix(0).miss.map(x => x.month + ' ' + x.post); state.rosters = keep; return r; })()")
    ok(len(nr) > 100 and "2026-08 エジプト" in nr and "2026-08 ケニア" not in nr, f"配置表がないときの足りない月報 {len(nr)}件")
    app.tab("import")
    txt = app.page.inner_text("#submitCard")
    ok("足りない月報" in txt and "タイ" in txt and "未" in txt, "取り込みタブの提出状況")
    ok("ケニア" not in txt, f"そろっている公館は「足りない公館だけ」で隠れる: {txt}")
    app.page.uncheck("#sOnly")
    ok("ケニア" in app.page.inner_text("#submitCard"), "すべての公館の表示")
    app.page.check("#sOnly")
    app.tab("settings")
    rows = app.js("[...document.querySelectorAll('table.regdata tbody tr')].map(tr => tr.cells[0].textContent + '|' + tr.cells[3].textContent)")
    r8 = next(r for r in rows if r.startswith("2026年8月|"))
    r7 = next(r for r in rows if r.startswith("2026年7月|"))
    ok("タイ" in r8 and "ケニア" not in r8 and "フィリピン" not in r8 and "フィリピン" in r7 and "スーダン（退避中）" in r8 and "エジプト" not in r8, f"登録済みデータの未提出公館（配置表の公館だけ）: {r8[:80]}")


@test("配置表：番号つきの配置表PDFの文字（公館名が行から外れる形）を読み取れる")
def t_roster_numbered(app):
    # 架空の名前。結合したセルの公館（タイ・アメリカ合衆国・南アフリカ共和国）は行から外れて別の場所に出る
    txt = "\n".join([
        "駐在公館名広産氏　　　名ふりがな発令年月日専　     門",
        "1 インド架空　一郎かくう　いちろうR07.01.24 救急医学",
        "2 インドネシア見本　二郎みほん　じろうR07.03.30 内科",
        "3 スリランカ産試験　花子しけん　はなこR08.07.01 精神科",
        "4 広検証　三郎けんしょう　さぶろうR08.08.31 外科",
        "5 メ試作　四郎しさく　しろうR07.03.16 精神科",
        "6 中華人民共和国産仮名　五郎かめい　ごろうR06.07.15 内科",
        "7 広州総見本　六子みほん　むつこR08.04.01 外科",
        "8 ミクロネシア",
        "9 広産架空　七子かくう　ななこR06.07.05 内科",
        "10 メ",
        "11 ニューヨーク総産試験　八郎しけん　はちろうR08.07.20 麻酔科",
        "12 キューバ遠藤　ひろみえんどう　ひろみR08.03.02 内科",
        "13 トリニダード・トバゴ未配置",
        "14 フランス広産仮名　九郎かめい　くろうR07.01.31 内科",
        "15 メ産検証　十子けんしょう　とおこR07.03.30 精神科",
        "16 ユジノサハリンスク総",
        "17 アフガニスタンアメリカ合衆国医務官等配置表",
        "(令和８年１０月１日現在 在外公館定員１４公館１９名 (退避１公館、空席４公館、未配置１公館）)",
        "タイ",
        "18 イエメン退避中",
        "19 マリ見本　花江みほん　はなえR08.06.22 内科",
        "1 インド看護　師子かんご　しこR06.09.20 看護学",
        "出　向南アフリカ共和国",
    ])
    r = app.js("(t) => { const o = parseRosterText(t); return {rows: o.map(x => [x.post, x.officer, x.note || '', x.staffed]), month: o.sheetMonth, cap: o.capacity}; }", txt)
    rows = {x[0]: x for x in r["rows"]}
    eq(r["month"], "2026-10", "見出しの年月")
    eq(r["cap"], {"posts": 14, "slots": 19}, "定員")
    eq(sorted(rows), sorted(["インド", "インドネシア", "スリランカ", "タイ", "中国", "広州総領事館", "ミクロネシア", "アメリカ合衆国", "ニューヨーク総領事館",
        "キューバ", "トリニダード・トバゴ", "フランス", "ユジノサハリンスク総領事館", "アフガニスタン", "イエメン", "マリ"]), "読み取った公館")
    eq(rows["タイ"][1:], ["検証 三郎", "", True], "行から外れた公館名（タイ）")
    eq(rows["アメリカ合衆国"][1:], ["架空 七子", "", True], "行から外れた公館名（米国）")
    eq(rows["ミクロネシア"][1:], ["", "空席", False], "空席の公館")
    eq(rows["キューバ"][1], "遠藤 ひろみ", "名がひらがなの人")
    eq(rows["フランス"][1], "仮名 九郎", "1公館2人")
    eq(app.js("(t) => parseRosterText(t).filter(x => x.officers.length > 1).map(x => x.post + ':' + x.officers.map(o => o.mental ? 'メ' : '身').join(''))", txt),
       ["タイ:身メ", "アメリカ合衆国:身メ", "フランス:身メ"], "身体とメンタルの2人体制")
    eq([rows["トリニダード・トバゴ"][2], rows["イエメン"][2], rows["アフガニスタン"][2]], ["未配置", "退避中", "空席"], "未配置・退避中・空席")
    ok(rows["インド"][1] == "架空 一郎", "看護師の欄（番号が1に戻る）を医務官として読んだ")


@test("医務官の異動：ケニア → ウガンダ を検出")
def t_transfer(app):
    t = app.js("transfers().filter(x => x.month === '2026-08').map(x => [x.post, x.before, x.after, x.from])")
    ok(["ウガンダ", "試験 次郎", "架空 一郎", "ケニア"] in t, f"ウガンダの異動がない: {t}")
    ok(["ケニア", "架空 一郎", "見本 花子", None] in t, f"ケニアの交代がない: {t}")


@test("巡回健診・感染症・緊急移送が保存される")
def t_events(app):
    eq(app.js("state.junkai['2026-08'].map(j => [j.from, j.dest])"), [["ウガンダ", "南スーダン"]], "8月の巡回健診")
    ob = app.js("outbreaksOf('2026-08').map(o => [o.post, o.name, o.trend]).sort()")
    eq(ob, [["ケニア", "コレラ", "減少・終息"], ["フィリピン", "デング熱", "流行・増加"]], "8月の感染症")
    eq(app.js("casesOf('2026-07', 'evac').map(c => [c.post, c.dest])"), [["ケニア", "南アフリカ共和国"]], "7月の緊急移送")


@test("10行サマリー：医療事情を厚く（4行）、見本の内容どおりに作る／以前の自動作成は作り直す")
def t_summary(app):
    lines = app.js("ruleSummary('2026-08', 'ケニア')").split("\n")
    eq(len(lines), 10, "行数")
    eq(lines[0], "受診者数：20名（職員12名、家族6名、その他2名）で、前月（18名）から2名（11%）増加した。", "1行目")
    ok(lines[1].startswith("主な疾病：") and "6件" in lines[1] and lines[1].endswith("。"), f"2行目: {lines[1]}")
    eq(lines[3], "感染症の動向：コレラは減少・終息に向かっている。", "4行目")
    eq(lines[4], "感染症・公衆衛生：コレラの新規報告は減少傾向にある。", "5行目")
    ok("着任" in lines[8], f"9行目: {lines[8]}")
    ok(app.js("ruleSummary('2026-08', 'フィリピン').split('\\n')[0]").endswith("であった（前月の月報なし）。"), "前月がないとき")
    eq(app.js("ruleSummary('2026-07', 'ケニア').split('\\n')[7]"), "邦人保護・緊急対応：邦人旅行者1名が交通外傷で入院し、南アフリカの病院へ緊急移送した。手術を受け回復した。", "邦人保護の行（経過の文も続けて）")
    eq(app.js("ruleSummary('2026-07', 'ケニア').split('\\n')[5]"), "医療機関：特記なし", "事例の経過の文を医療機関の行に使わない")
    # 医療の記述が多い丙表：感染症・医療機関・医薬品がそれぞれの行に入り、キーワードのない医療の文も空き行を埋める
    hei = ("1. 医療事情\n市内の私立病院でICUの病床がひっ迫し、救急の受け入れを制限している。"
           "インスリンなど一部の医薬品が品薄で、館員に常備薬の確保を周知した。\n"
           "2. 感染症\n雨季に入りデング熱が急増、保健省が注意喚起を出した。\n3. 治安\n首都でデモが続いている。")
    med = app.js("""(h) => { state.reports['2026-09|ケニア'] = {month: '2026-09', post: 'ケニア', region: regionOf('ケニア'), total: 5, staff: 5, family: 0, other: 0, diseases: [], hei: h};
        const t = ruleSummary('2026-09', 'ケニア'); delete state.reports['2026-09|ケニア']; return t.split('\\n'); }""", hei)
    ok("デング熱" in med[4], f"感染症・公衆衛生: {med[4]}")
    ok("ICU" in med[5], f"医療機関: {med[5]}")
    ok("医薬品" in med[6], f"医薬品・医療体制: {med[6]}")
    ok("デモ" in med[9], f"治安: {med[9]}")
    hei2 = "1. 医療事情\n公立の医療は無償だが待ち時間が長い状況が続いている。\n2. その他\n特記事項なし。"
    med2 = app.js("""(h) => { state.reports['2026-09|ケニア'] = {month: '2026-09', post: 'ケニア', region: regionOf('ケニア'), total: 5, staff: 5, family: 0, other: 0, diseases: [], hei: h};
        const t = ruleSummary('2026-09', 'ケニア'); delete state.reports['2026-09|ケニア']; return t.split('\\n'); }""", hei2)
    ok("待ち時間" in "".join(med2[4:7]), f"医療の見出しの下の文が載らない: {med2[4:7]}")
    # 以前の自動作成（5行・10行）は作り直し、手で書いたものは残す
    # 自然な文：文末は「。」、矢印や「(件数)」の書き方を使わない
    for l in app.js("ruleSummary('2026-07', 'ケニア', 15)").split("\n"):
        body = l.split("：", 1)[1]
        ok(body == "特記なし" or body.endswith("。") or body.endswith("…"), f"文末が「。」でない: {l}")
        ok("↑" not in l and "↓" not in l and "→" not in l and not re.search(r"\(\d+\)", l), f"記号が残っている: {l}")
    ok("、" in app.js("ruleSummary('2026-07', 'ケニア', 15)").split("\n")[12] and "緊急移送し" in app.js("ruleSummary('2026-07', 'ケニア', 15)").split("\n")[12], "緊急移送の文")
    old5 = "受診者数 18名（職員10・家族5・その他3）\n主な疾病：かぜ(4)\n医療・感染症：特記なし\n邦人保護・緊急対応：特記なし\nその他：特記なし"
    old10 = "\n".join(f"{k}：x" for k in ["受診者数", "前月比", "主な疾病", "疾病分類", "検診・予防接種", "感染症の動向", "医療事情", "邦人保護・緊急対応", "巡回健診・医務官", "治安・災害・その他"])
    app.js("([a, b]) => { state.reports['2026-07|ケニア'].summary = a; state.reports['2026-08|ケニア'].summary = b; const t = state.reports['2026-07|タイ']; t.summary = '手で書いた要約'; t.summaryManual = true; }", [old5, old10])
    eq(app.js("upgradeSummaries()") >= 2, True, "作り直した件数")
    eq(app.js("state.reports['2026-07|ケニア'].summary"), app.js("ruleSummary('2026-07', 'ケニア')"), "5行の作り直し")
    eq(app.js("state.reports['2026-08|ケニア'].summary"), app.js("ruleSummary('2026-08', 'ケニア')"), "以前の10行の作り直し")
    eq(app.js("state.reports['2026-07|タイ'].summary"), "手で書いた要約", "手で書いた要約")
    app.js("const t = state.reports['2026-07|タイ']; t.summary = ruleSummary('2026-07', 'タイ'); t.summaryManual = false;")


@test("10行サマリー：「自動作成し直す」で書き換えた文が戻り、押したことが分かる")
def t_regen(app):
    app.js("refreshMonths('2026-08')")
    app.tab("summary")
    app.js("window._toasts = []; const _t = toast; toast = (m, ...a) => { _toasts.push(String(m)); return _t(m, ...a); }")
    ta = app.page.locator(".sumcard textarea").first
    auto = ta.input_value()
    ta.fill("手で書いた文")
    # 今の文が消えるときは確認が出る（やめると何も変わらない）
    msgs = []
    app.page.once("dialog", lambda d: (msgs.append(d.message), d.dismiss()))
    app.page.locator(".sumcard [data-act=regen]").first.click()
    app.page.wait_for_timeout(300)
    ok(msgs and "消えて" in msgs[0], "書き直す前の確認が出ない")
    eq(ta.input_value(), "手で書いた文", "やめたのに書き直された")
    app.page.once("dialog", lambda d: d.accept())
    app.page.locator(".sumcard [data-act=regen]").first.click()
    app.page.wait_for_function("_toasts.some(m => m.includes('自動作成し直しました'))")
    eq(ta.input_value(), auto, "作り直した文")
    app.page.locator(".sumcard [data-act=regen]").first.click()
    app.page.wait_for_function("_toasts.some(m => m.includes('今の文面と同じ'))")
    app.page.once("dialog", lambda d: d.accept())
    app.page.click("#regenAll")
    app.page.wait_for_function("_toasts.some(m => m.includes('公館のサマリーを自動作成し直しました'))")


@test("サマリー：5行・10行・15行をその場で切り替え（手で書いたものは残す・PDFはPDFレポートで選んだ行数。はじめは5行）")
def t_sizes(app):
    app.js("refreshMonths('2026-08')")
    app.tab("summary")
    app.js("window._toasts = []; const _t = toast; toast = (m, ...a) => { _toasts.push(String(m)); return _t(m, ...a); }")
    app.js("const t = state.reports['2026-08|フィリピン']; t.summary = '手で書いた要約'; t.summaryManual = true;")
    for n in [5, 15, 10]:
        app.page.click(f'[data-slines="{n}"]')
        app.page.wait_for_function(f"_toasts.some(m => m.includes('{n}行にしました'))")
        eq(app.js("document.querySelector('main h2').textContent").startswith(f"公館別 {n}行サマリー"), True, f"{n}行の見出し")
        k = app.js("[...document.querySelectorAll('.sumcard')].findIndex(c => c.dataset.post === 'ケニア')")
        eq(len(app.page.locator(".sumcard textarea").nth(k).input_value().split("\n")), n, f"{n}行のときのケニア")
        eq(app.js("state.reports['2026-08|フィリピン'].summary"), "手で書いた要約", f"{n}行に変えても手で書いたものは残る")
        # PDFは画面の行数にかかわらず、PDFレポートで選んだ行数（はじめは5行。手で書いたものはその文のまま）
        text5 = app.js("async () => (await reportBlocks('2026-08', new Set(['summary']), {label: '', scope: '', sumLines: 5})).map(e => e.textContent).join('\\n')")
        ok("公館別 5行サマリー" in text5 and "受診者数の推移" not in text5, f"{n}行のときPDFが5行にならない")
        text = app.js("async () => (await reportBlocks('2026-08', new Set(['summary']), {label: '', scope: '', sumLines: 15})).map(e => e.textContent).join('\\n')")
        ok("公館別 15行サマリー" in text and "受診者数の推移" in text, f"PDFで15行を選んだのに15行にならない")
        ok("特記なし" not in text, "PDFのサマリーに「特記なし」の行が載っている")
        ok("手で書いた要約" in text, "PDFに手で書いた要約が載らない")
    five = app.js("ruleSummary('2026-07', 'ケニア', 5)").split("\n")
    eq([l.split("：")[0] for l in five], ["受診者数", "主な疾病", "医療・感染症", "邦人保護・緊急対応", "その他"], "5行の項目")
    ok("コレラ" in five[2], f"5行の医療・感染症: {five[2]}")
    ft = app.js("ruleSummary('2026-08', 'ケニア', 15)").split("\n")
    eq(ft[10], "受診者数の推移：7月18名、8月20名と推移した。", "15行の推移")
    ok(app.js("ruleSummary('2026-07', 'ケニア', 15)").split("\n")[12].startswith("緊急移送：1件"), "15行の緊急移送")
    app.js("const t = state.reports['2026-08|フィリピン']; t.summary = ruleSummary('2026-08', 'フィリピン'); t.summaryManual = false;")


@test("サマリー：GaiXer／源内に渡す文章を作り、AIの回答を数値を確かめて取り込む")
def t_ai_summary(app):
    app.js("refreshMonths('2026-07')")
    app.tab("summary")
    app.js("state.reports['2026-07|ケニア']._hei = state.reports['2026-07|ケニア'].hei; state.reports['2026-07|ケニア'].hei += '\\n35歳の男性館員が受診した。連絡先 +254-20-1234567。'")
    ok(not app.js("document.querySelector('#aisum').open"), "最初から手順の欄が開いている")
    app.page.click("#aisStart")
    ok(app.js("document.querySelector('#aisum').open"), "「AIでサマリーを作る」で手順の欄が開かない")
    t = app.page.input_value("#aisText")
    ok(t.startswith("【取扱注意】"), "冒頭の注意書き")
    for w in ["【ケニア】", "【タイ】", "【ウガンダ】", "乙表（病名と件数）", "丙表（報告事項）", "数値は公館データにあるものだけ"]:
        ok(w in t, f"AIに渡す文章に「{w}」がない")
    for w in ["35歳", "男性", "+254-20-1234567", "架空 一郎", "試作 七美"]:
        ok(w not in t, f"AIに渡す文章に個人情報「{w}」が残っている")
    # 個人情報の指摘：箇所を前後の文つきで示し、1件ずつ「この箇所へ」「伏せる」「問題なし」で片付ける
    orig = app.page.input_value("#aisText")
    app.page.fill("#aisText", orig + "\n山田様が受診した。35歳の男性。連絡は +81-3-1234-5678 へ。架空 一郎医務官が対応。")
    items = app.js("[...document.querySelectorAll('#aisCheck [data-pv]')].map(li => [li.querySelector('.badge').textContent, li.querySelector('mark').textContent])")
    eq(items, [["敬称（個人名の可能性）", "山田様"], ["年齢", "35歳"], ["性別・続柄", "男性"], ["電話番号", "+81-3-1234-5678"], ["医務官名", "架空 一郎"]], "指摘の箇所")
    app.page.click('#aisCheck [data-pv-go="1"]')
    eq(app.js("(() => { const t = document.querySelector('#aisText'); return t.value.slice(t.selectionStart, t.selectionEnd); })()"), "35歳", "「この箇所へ」で文章欄のその場所を選ぶ")
    app.page.click('#aisCheck [data-pv-fix="0"]')
    ok("（氏名）が受診した" in app.page.input_value("#aisText"), "「伏せる」で置き換わらない")
    eq(app.js("document.querySelectorAll('#aisCheck [data-pv]').length"), 4, "伏せた後の残り")
    app.page.click('#aisCheck [data-pv-ok="0"]')
    eq(app.js("document.querySelectorAll('#aisCheck [data-pv]').length"), 3, "「問題なし」の後の残り")
    ok(app.page.is_disabled('[data-ais-go="gaixer"]'), "指摘が残っているのに送れる")
    # 「すべて確認した・問題なし」：確認の上で、残りをまとめて問題なしにする（文章は書き換えない）
    before = app.page.input_value("#aisText")
    app.page.once("dialog", lambda d: d.accept())
    app.page.click("#aisCheck [data-pv-okall]")
    ok("✓" in app.page.inner_text("#aisCheck") and app.page.input_value("#aisText") == before, "すべて問題なしにできない")
    app.page.fill("#aisText", before + "\n田中さんが受診。")
    app.page.click("#aisCheck [data-pv-fixall]")
    ok("✓" in app.page.inner_text("#aisCheck"), "すべて伏せた後に点検が通らない")
    app.page.fill("#aisText", orig)
    ok(app.page.is_disabled('[data-ais-go="gaixer"]'), "確認のチェック前に送れる")
    app.page.check("#aisOk")
    ok(not app.page.is_disabled('[data-ais-go="gaixer"]'), "確認後に送れない")
    answer = ("※この集計は外務省AI（MOFA GaiXer）・政府AI源内専用です。\n\n"
              "【ケニア】\n受診者数：18名（職員10名、家族5名、その他3名）であった。\n主な疾病：かぜ4件、胃腸炎3件が多かった。\n\n"
              "**【タイ】**\n- 受診者数：10名であった。\n- 主な疾病：便秘が99件であった。\n\n【架空国】\n受診者数：1名")
    app.page.fill("#aisAnswer", answer)
    app.page.click("#aisParse")
    rows = app.js("[...document.querySelectorAll('#aisPreview tr')].slice(1).map(tr => tr.innerText)")
    eq(len(rows), 2, "読み取った公館の数")
    ok("数値OK" in rows[0] and "要確認" in rows[1] and "99" in rows[1], f"数値の照合: {rows}")
    ok("架空国" in app.page.inner_text("#aisPreview"), "月報にない公館名の知らせ")
    app.page.uncheck('[data-ais-i="1"]')
    app.page.click("#aisApply")
    app.page.wait_for_function("!!state.reports['2026-07|ケニア'].summaryAI")
    eq(app.js("state.reports['2026-07|ケニア'].summary"), "受診者数：18名（職員10名、家族5名、その他3名）であった。\n主な疾病：かぜ4件、胃腸炎3件が多かった。", "取り込んだサマリー")
    eq(app.js("!!state.reports['2026-07|タイ'].summaryAI"), False, "選ばなかった公館を取り込んだ")
    app.page.wait_for_function("document.querySelector('#slist').innerText.includes('AI作成')", timeout=5000)
    eq(app.js("upgradeSummaries() >= 0 && state.reports['2026-07|ケニア'].summary.startsWith('受診者数：18名（職員10名')"), True, "AI作成のサマリーが自動作成で上書きされた")
    pdf = app.js("async () => (await reportBlocks('2026-07', new Set(['summary']))).map(e => e.textContent).join('\\n')")
    ok("かぜ4件、胃腸炎3件が多かった" in pdf, "PDFにAI作成のサマリーが載らない")
    k = app.js("[...document.querySelectorAll('.sumcard')].findIndex(c => c.dataset.post === 'ケニア')")
    app.page.once("dialog", lambda d: d.accept())
    app.page.locator(".sumcard [data-act=regen]").nth(k).click()
    app.page.wait_for_function("!state.reports['2026-07|ケニア'].summaryAI")
    app.js("const r = state.reports['2026-07|ケニア']; r.hei = r._hei; delete r._hei; r.summary = ruleSummary('2026-07', 'ケニア');")


@test("統計：自由集計で好きなグラフ・記述統計・集計表を作れる")
def t_free_stats(app):
    app.js("refreshMonths('2026-08')")
    app.tab("stats")
    app.page.wait_for_selector("#fsCard")
    def setv(**kw):
        for k, v in kw.items(): app.page.select_option(f"#fs{k}", str(v))
        app.page.wait_for_timeout(200)
    setv(Span=0, Measure="total", Agg="sum", X="month", Color="", Chart="bar")
    rows = app.js("[...document.querySelectorAll('#fsStats table')[1].querySelectorAll('tbody tr')].map(tr => [...tr.cells].map(c => c.textContent))")
    eq(rows, [["2026年7月", "35"], ["2026年8月", "46"], ["合計", "81"]], "月ごとの受診者数の合計")
    setv(X="region", Chart="box")
    d = app.js("[...document.querySelectorAll('#fsStats table')[0].querySelectorAll('tbody tr')].map(tr => [tr.cells[0].textContent, tr.cells[1].textContent, tr.cells[2].textContent])")
    eq(d, [["アフリカ", "4", "56"], ["アジア", "2", "25"], ["全体", "6", "81"]], "地域ごとの記述統計（件数・合計）")
    setv(X="kind", Chart="pie")
    ok("受診者区分で分けたときの指標は受診者数" in app.page.inner_text("#fsMsg"), "受診者区分の説明")
    setv(X="month", Color="cat", Chart="stack100")
    ok(app.js("document.querySelector('#fsPlot').data.length") >= 3, "疾病分類の100%積み上げ")
    setv(Chart="scatter", Measure="total", Measure2="real", X="post", Color="")
    ok("ピアソンの相関係数" in app.page.inner_text("#fsStats"), "散布図の相関")
    app.page.click('[data-fsreg="アフリカ"]')
    ok("月報4件" in app.page.inner_text("#fsMsg"), "地域で絞る")
    app.page.click('[data-fsreg=""]')
    app.page.click('[data-fsp="0"]')
    eq(app.page.input_value("#fsChart"), "line", "よく使う例")
    eq(app.errors, [], "画面のエラー")


@test("統計：検定・推定（計算は SciPy と一致し、結果をレポートに入れられる）")
def t_inference(app):
    # 計算の正しさ（SciPy で求めた値と比べる）
    r = app.js("""(() => { const a = [12,15,9,20,17,14,11,13,19,16], b = [22,18,25,19,21,27,23,20,24], c = [14,16,15,18,13,17,19];
        return [ST.welch(a, b).t, ST.welch(a, b).p, ST.mannWhitney(a, b).p, ST.anova([a, b, c]).F, ST.kruskal([a, b, c]).H, ST.paired(a.slice(0, 9), b).p,
                ST.chisq([[30,12,8],[20,25,15],[10,9,21]]).x2, ST.poissonUpper(12, 5.3), ST.tQuantile(0.975, 7.3)]; })()""")
    ref = [-5.08313783, 9.329e-05, 0.00080644, 16.14731289, 14.79556819, 0.00299977, 22.61157773, 0.00840873, 2.34506674]
    for got, want in zip(r, ref):
        ok(abs(got - want) <= 1e-6 * max(1, abs(want)), f"計算が SciPy と違う: {got} / {want}")
    app.js("refreshMonths('2026-08')")
    app.tab("stats")
    app.page.wait_for_selector("#infCard")
    app.page.select_option("#infSpan", "0"); app.page.select_option("#infType", "two")
    app.page.select_option("#infA", "アフリカ"); app.page.select_option("#infB", "アジア"); app.page.select_option("#infMeasure", "total")
    t = app.page.inner_text("#infOut")
    ok("アフリカの平均は14" in t and "アジアの平均は12.5" in t and "ウェルチのt検定" in t, f"2つの地域の比較: {t[:200]}")
    for typ, w in [("multi", "一元配置分散分析"), ("prev", "前月との比較"), ("chi", "カイ二乗検定"), ("prop", "割合の差"), ("excess", "通り（公館×病名）"), ("corr", "相関"), ("trend", "4か月以上のデータが必要")]:
        app.page.select_option("#infType", typ)
        if typ == "prev": app.page.select_option("#infSpan", "1")
        if typ == "corr": app.page.select_option("#infSpan", "0")  # 相関は4公館以上が必要（8月だけだと3公館）
        ok(w in app.page.inner_text("#infOut"), f"「{typ}」の結果に「{w}」がない: {app.page.inner_text('#infOut')[:160]}")
    eq(app.errors, [], "画面のエラー")
    # 保存してレポートへ
    app.page.select_option("#infType", "two"); app.page.select_option("#infSpan", "0")
    app.page.click("#infSave")
    app.page.wait_for_function("(state.statNotes || []).length === 1")
    rep = app.js("async () => (await reportBlocks('2026-08', new Set(['stats']), {label: reportLabel('2026-08'), scope: '全世界', stats: ['infer']})).map(e => e.textContent).join('\\n')")
    ok("検定・推定の結果" in rep and "アフリカ と アジア の比較" in rep and "ウェルチのt検定" in rep, "保存した結果がレポートに入らない")
    app.js("state.statNotes = []")


@test("再読み込みしてもデータが残る（ブラウザ内に保存）")
def t_persist(app):
    app.open()
    eq(app.js("Object.keys(state.reports).length"), 6, "再読み込み後の月報の件数")
    eq(app.js("Object.keys(state.rosters)"), ["2026-08"], "再読み込み後の配置表")


@test("2つのタブで開いたとき：片方で保存すると、もう片方も読み直す")
def t_two_tabs(app):
    # 新しく開いたタブは、すでに開いているタブがあれば自分を閉じる（タブが増えないように）
    dup = app.ctx.new_page()
    dup.goto(app.base + "/geppou/")
    dup.wait_for_timeout(2000)
    ok(dup.is_closed() or "すでに開いています" in dup.inner_text("body"), "2つ目のタブが閉じない・案内が出ない")
    if not dup.is_closed(): dup.close()
    # 「このタブで使う」を選んだタブは、そのまま使える
    other = app.ctx.new_page()
    other.add_init_script("try { sessionStorage.setItem('geppou.useThisTab', '1'); } catch (e) {}")
    other.goto(app.base + "/geppou/")
    other.wait_for_function("typeof state !== 'undefined' && Object.keys(state.reports).length === 6")
    app.js("state.groups['テスト病'] = 'テストまとめ'; saveState()")
    other.wait_for_function("state.groups['テスト病'] === 'テストまとめ'", timeout=10000)
    other.evaluate("delete state.groups['テスト病']; saveState()")
    app.page.wait_for_function("!state.groups['テスト病']", timeout=10000)
    other.close()
    eq(app.js("Object.keys(state.reports).length"), 6, "月報の件数")
    eq(app.js("[document.querySelector('#installBtn').hidden, document.querySelector('link[rel=manifest]') && 1]"), [False, 1], "デスクトップに置くボタン")


@test("すべてのタブがエラーなく表示できる")
def t_tabs(app):
    app.js("refreshMonths('2026-08')")
    for name in app.js("[...document.querySelectorAll('#tabs button[data-page]')].map(b => b.dataset.page)"):
        before = len(app.errors)
        app.tab(name)
        eq(app.errors[before:], [], f"「{name}」タブのエラー")
        ok(len(app.js("document.querySelector('main').innerText")) > 20, f"「{name}」タブが空")


# ----------------------------------------------------------------------- AI用テキスト・出力

@test("AI用テキスト：注意書き・管理番号が付き、氏名が入らない")
def t_ai(app):
    t = app.js("aiText('2026-08', new Set(['submit','visits','cats','top','events','hei','cases','quality']))")
    s = app.js("async (t) => { const n = (state.aiLog || []).length; const r = await stampAiText(t, 'テスト', {}); return [r.t, r.id, state.aiLog.length - n, state.aiLog[0].fp]; }", t)
    text, aid, added, fp = s
    ok(re.fullmatch(r"AI-\d{8}-[A-HJ-NP-Z2-9]{4}", aid), f"管理番号の形式: {aid}")
    ok(text.startswith("【取扱注意】この集計は外務省AI（MOFA GaiXer）・政府AI源内専用です。"), "冒頭の注意書き")
    ok(f"管理番号：{aid}" in text, "本文中の管理番号")
    ok("回答の最初に、次の注意文をそのまま表示" in text, "AIへの注意文の指示")
    eq(added, 1, "記録（aiLog）に1件追加")
    eq(app.js("(t) => textFingerprint(t.replace(/^管理番号：.*$/m, '管理番号：消された'))", text), fp, "番号を消しても指紋で照合できる")
    names = {o for _, o, *_ in fixtures.REPORTS.values()} | {n for _, n, _ in fixtures.ROSTER if n}
    for nm in names:
        ok(nm not in text and nm.replace(" ", "") not in text, f"氏名「{nm}」が入っている")
    ok("南アフリカ" not in app.js("(t) => privacyIssues(t).join()", text), "点検が誤検出していない")
    ok(app.js("privacyIssues('35歳の男性が受診').length") >= 2, "年齢・性別を点検で見つける")
    for t in ["様式Aを提出した。", "前月と同様に推移。", "雨模様が続いた。", "多様な症例。", "氏名は記載しない。", "霰粒腫（さん粒）1件。", "さん粒腫。", "たくさんの受診があった。", "お子さんの発熱。", "患者さんが増えた。"]:
        eq(app.js("(t) => privacyIssues(t)", t), [], f"「{t}」を敬称と誤検出")
    for t in ["山田様が受診。", "田中さんが入院。", "佐藤氏に連絡。"]:
        ok(any("敬称" in x for x in app.js("(t) => privacyIssues(t)", t)), f"「{t}」の敬称を見落とす")
    eq(app.js("aiUrls()"), {"gaixer": "https://mofa-ai.gaixer.com/general", "gennai": "https://d3j5pbzknc8vap.cloudfront.net/"}, "AIの送り先（既定）")


@test("設定：疾病分類をAIで決める（書き出し・回答の読み込み）／取り込み画面にパスワードを出さない")
def t_ai_classify(app):
    app.js("state.reports['2026-07|ケニア'].diseases.push({name: 'ほげほげ症', count: 2}); state.reports['2026-07|ケニア'].diseases.push({name: 'ナゾ熱', count: 1}); clearCaches();")
    app.tab("settings")
    t = app.js("clsPrompt(['ほげほげ症', 'ナゾ熱'], false)")
    ok(t.startswith("【取扱注意】") and "病名\t分類\t代表名" in t and "ほげほげ症" in t and "呼吸器系の疾患" in t, "AIに渡す依頼文")
    ok("ケニア" not in t, "依頼文に公館名が入っている")
    ans = "※この集計は外務省AI（MOFA GaiXer）・政府AI源内専用です。\n病名\t分類\t代表名\nほげほげ症\t呼吸器系の疾患\tほげ症候群\nナゾ熱（いまの分類：分類不能）\t感染症\t\n存在しない病名\t消化器系の疾患\t\nかぜ | 呼吸器"
    app.page.fill("#clsAns", ans)
    app.page.click("#clsParse")
    rows = app.js("[...document.querySelectorAll('#clsPreview tbody tr')].map(tr => [tr.cells[1].textContent, tr.cells[3].textContent, tr.querySelector('input').checked])")
    eq(rows[:2], [["ほげほげ症", "呼吸器", True], ["ナゾ熱", "感染症", True]], "回答の読み取り")
    ok("存在しない病名" in app.page.inner_text("#clsPreview"), "月報にない病名の知らせ")
    app.page.click("#clsApply")
    app.page.wait_for_function("state.overrides['ほげほげ症'] === '呼吸器系の疾患'")
    eq(app.js("[classify('ナゾ熱'), groupOf('ほげほげ症')]"), ["感染症及び寄生虫症", "ほげ症候群"], "保存した分類・まとめ")
    app.js("delete state.overrides['ほげほげ症']; delete state.overrides['ナゾ熱']; delete state.groups['ほげほげ症']; for (const k of ['2026-07|ケニア']) state.reports[k].diseases = state.reports[k].diseases.filter(d => !['ほげほげ症','ナゾ熱'].includes(d.name)); clearCaches();")
    app.tab("import")
    ok("imukanmd" not in app.js("document.querySelector('main').innerText"), "取り込み画面にパスワードが出ている")


@test("AIでまとめる：病名の一覧を書き出し、AIの回答（病名⇥まとめ先）を読み込んで保存できる")
def t_ai_group(app):
    app.js("state.reports['2026-07|ケニア'].diseases.push({name: 'ほげほげ症', count: 2}, {name: 'ホゲホゲ病', count: 1}); clearCaches();")
    app.tab("settings")
    t = app.js("grpPrompt(['ほげほげ症', 'ホゲホゲ病'], false)")
    ok(t.startswith("【取扱注意】") and "病名\tまとめ先" in t and "ホゲホゲ病" in t and "デング熱" in t, "AIに渡す依頼文")
    ok("ケニア" not in t, "依頼文に公館名が入っている")
    ans = "病名\tまとめ先\nほげほげ症\tほげ症候群\nホゲホゲ病（いまのまとめ先：ホゲホゲ病）\tほげ症候群\n存在しない病名\tなにか\n急性鼻咽頭炎 → かぜ（急性上気道炎）"
    app.page.click("#aigrp summary")
    app.page.fill("#grpAns", ans)
    app.page.click("#grpParse")
    rows = app.js("[...document.querySelectorAll('#grpPreview tbody tr')].map(tr => [tr.cells[1].textContent, tr.cells[4].textContent, tr.querySelector('input').checked])")
    ok(["ほげほげ症", "ほげ症候群 新しいまとめ先", True] in rows and ["ホゲホゲ病", "ほげ症候群 新しいまとめ先", True] in rows, f"回答の読み取り {rows}")
    ok("存在しない病名" in app.page.inner_text("#grpPreview"), "月報にない病名の知らせ")
    app.page.click("#grpApply")
    app.page.wait_for_function("state.groups['ホゲホゲ病'] === 'ほげ症候群'")
    eq(app.js("[groupOf('ほげほげ症'), groupOf('ホゲホゲ病')]"), ["ほげ症候群", "ほげ症候群"], "保存したまとめ先")
    app.js("delete state.groups['ほげほげ症']; delete state.groups['ホゲホゲ病']; state.reports['2026-07|ケニア'].diseases = state.reports['2026-07|ケニア'].diseases.filter(d => !['ほげほげ症','ホゲホゲ病'].includes(d.name)); clearCaches();")


@test("設定：チェックした病名を分類済みにして一覧から消し、「分類済みを表示」で見られる")
def t_done_rows(app):
    app.js("state.reports['2026-07|ケニア'].diseases.push({name: 'ほげほげ症', count: 2}, {name: 'ナゾ熱', count: 1}); clearCaches();")
    app.tab("settings")
    ok(app.js("!!document.querySelector('[data-ovchk=\"ほげほげ症\"]')"), "疾病分類の表にない")
    app.page.select_option('select[data-name="ほげほげ症"]', label="感染症")
    app.page.check('[data-ovchk="ほげほげ症"]')
    app.page.click("#ovHide")
    app.page.wait_for_function("!document.querySelector('[data-ovchk=\"ほげほげ症\"]')")
    eq(app.js("[state.doneCls['ほげほげ症'], classify('ほげほげ症')]"), [True, "感染症及び寄生虫症"], "分類済みにした病名の分類")
    ok(app.js("!!document.querySelector('[data-ovchk=\"ナゾ熱\"]')"), "チェックしていない病名まで消えた")
    app.page.click("#ovShowDone")
    app.page.wait_for_function("!!document.querySelector('[data-ovchk=\"ほげほげ症\"]')")
    app.page.check('[data-ovchk="ほげほげ症"]')
    app.page.click("#ovUnhide")
    app.page.wait_for_function("(() => { const c = document.querySelector('[data-ovchk=\"ほげほげ症\"]'); return c && !c.closest('tr').classList.contains('done'); })()")
    # 病名のまとめ方
    app.page.fill('[data-g="ナゾ熱"]', "なぞの熱")
    app.page.check('[data-gchk="ナゾ熱"]')
    app.page.click("#gHide")
    app.page.wait_for_function("!document.querySelector('[data-gchk=\"ナゾ熱\"]')")
    eq(app.js("[state.doneGrp['ナゾ熱'], groupOf('ナゾ熱')]"), [True, "なぞの熱"], "分類済みにした病名のまとめ先")
    ok("分類済みを表示（1件）" in app.page.inner_text("#gShowDone"), "分類済みを表示ボタン")
    app.page.click("#gShowDone")
    app.page.wait_for_function("!!document.querySelector('[data-gchk=\"ナゾ熱\"]')")
    app.js("delete state.overrides['ほげほげ症']; delete state.groups['ナゾ熱']; state.doneCls = {}; state.doneGrp = {}; showDoneGrp = showDoneCls = false; state.reports['2026-07|ケニア'].diseases = state.reports['2026-07|ケニア'].diseases.filter(d => !['ほげほげ症','ナゾ熱'].includes(d.name)); clearCaches(); saveState()")


@test("設定：疾病分類・病名のまとめ先をCSVで書き出せる")
def t_disease_csv(app):
    app.tab("settings")
    with app.page.expect_download() as dl:
        app.page.click("[data-dis-csv] >> nth=0")
    raw = open(dl.value.path(), "rb").read()
    ok(raw.startswith(b"\xef\xbb\xbf"), "BOM がない（Excel で文字化け）")
    lines = raw.decode("utf-8-sig").split("\r\n")
    eq(lines[0], '"病名","件数（全期間）","報告した公館数","疾病分類（ICD-10大分類）","分類の決め方","まとめ先（代表名）","まとめの決め方","分類済み"', "見出し")
    ok(any(l.startswith('"デング熱"') and "感染症" in l for l in lines), "デング熱の行")
    ok(dl.value.suggested_filename.startswith("疾病分類_病名一覧_") and dl.value.suggested_filename.endswith(".csv"), "ファイル名")
    ok(not any(p in raw.decode("utf-8-sig") for p in ["ケニア", "架空 一郎"]), "公館名・氏名が入っている")


@test("月報くんフォルダ：バックアップフォルダに日時入りで保存し、最新のものを読み込む")
def t_backup_folder(app):
    # フォルダを選ぶ画面は自動では押せないので、ブラウザ内の専用フォルダ（OPFS）を「バックアップ」フォルダの代わりに使う
    app.js("""async () => { const root = await navigator.storage.getDirectory(); const d = await root.getDirectoryHandle('バックアップ', { create: true });
        for await (const [n] of d.entries()) await d.removeEntry(n);
        bkdir.handle = d; bkdir.ok = true; }""")
    n = app.js("Object.keys(state.reports).length")
    app.js("window._keepState = JSON.parse(JSON.stringify(state))")
    app.js("backupToFolder(true)")
    names = app.js("async () => (await bkdirList()).map(x => x.name)")
    ok(len(names) == 1 and re.fullmatch(r"月報くん_\d{4}-\d{2}-\d{2}_\d{2}-\d{2}-\d{2}\.geppou", names[0]), f"日時入りのファイル名: {names}")
    ok(app.js("loadPrefs().lastBackupFile") == names[0] and ":" in app.js("backupStatusText()"), "最後のバックアップに時刻が出ない")
    # 別のパソコン・ブラウザで新しいバックアップができた（フォルダに新しいファイル）→ 開いたときに自動で読み込む
    app.js("""async () => { const s = JSON.parse(JSON.stringify(state)); delete s.reports['2026-07|ケニア'];
        const fh = await bkdir.handle.getFileHandle('月報くん_2099-01-01_09-00-00.geppou', { create: true }); const w = await fh.createWritable();
        await w.write(await encryptBackup(s)); await w.close(); savePref('lastChange', Date.now() - 1000); }""")
    app.js("bkdirSync()")
    eq(app.js("[Object.keys(state.reports).length, loadPrefs().lastBackupFile]"), [n - 1, "月報くん_2099-01-01_09-00-00.geppou"], "最新のバックアップを読み込む")
    # もう一度開いても、同じものは読み直さない
    app.js("state.reports['2026-07|ケニア'] = {month: '2026-07', post: 'ケニア', diseases: []}")
    app.js("bkdirSync()")
    eq(app.js("Object.keys(state.reports).length"), n, "同じバックアップを読み直した")
    # 「💾 バックアップ」を押すと、つないだフォルダに保存する（ダウンロードしない）
    k0 = app.js("async () => (await bkdirList()).length")
    app.page.wait_for_timeout(1100)   # ファイル名の秒が変わるように
    app.page.click("#backupTop")
    app.page.wait_for_timeout(800)
    eq(app.js("async () => (await bkdirList()).length"), k0 + 1, "バックアップのボタンでフォルダに保存")
    # 月報くんフォルダ（zip）
    app.tab("settings")
    with app.page.expect_download() as dl:
        app.page.click("#mkFolder")
    import zipfile
    zf = zipfile.ZipFile(dl.value.path())
    names = sorted(zf.namelist())
    ok("月報くん/月報くん.url" in names and any(x.startswith("月報くん/バックアップ/") for x in names) and "月報くん/はじめにお読みください.txt" in names, f"zip の中身: {names}")
    ok("URL=http" in zf.read("月報くん/月報くん.url").decode(), "ショートカットのアドレス")
    app.js("bkdir.handle = null; bkdir.ok = false; savePref('lastBackupFile', null); clearTimeout(bkdir.timer); state = _keepState; clearCaches(); saveState()")


@test("OCR：文字のないPDFを、文字認識で読み直せる（ブラウザの中だけで）")
def t_ocr(browser, base):
    app = App(browser, base)
    try:
        folder = tempfile.mkdtemp(prefix="geppou_ocr_")
        path = os.path.join(folder, "ケニア_202607_月報.pdf")
        fixtures.make_image_pdf(fixtures.pages_of("ケニア_202607_月報.pdf")[:1], path)
        app.tab("import")
        app.page.check('input[name="mmode"][value="auto"]')
        app.page.set_input_files("#files", [path])
        app.page.click("#start")
        app.page.wait_for_function("!importBusy && drafts.length === 1", timeout=60000)
        ok(app.js("drafts[0].total") is None and app.page.locator("[data-ocr]").count() == 1, "文字のないPDFにOCRのボタンが出ない")
        app.page.click("[data-ocr]")
        app.page.wait_for_function("drafts[0].ocr === true", timeout=240000)
        d = app.js("({post: drafts[0].post, officer: drafts[0].officer, counts: [drafts[0].staff, drafts[0].family, drafts[0].other, drafts[0].total], w: drafts[0].warnings})")
        ok(any("OCR" in w for w in d["w"]), "OCRで読んだ印がない")
        eq(d["counts"], [10, 5, 3, 18], f"OCRで読んだ受診者数 {d}")
        # ページ番号だけが文字で、本文が画像のPDF（ページ番号のある画像のPDF）も、文字のないページとして扱い、OCRで読める
        path2 = os.path.join(folder, "ケニア_202608_月報.pdf")
        fixtures.make_image_pdf(fixtures.pages_of("ケニア_202608_月報.pdf")[:1], path2, footer="- 1 -")
        app.tab("import")
        app.page.set_input_files("#files", [path2])
        app.page.click("#start")
        app.page.wait_for_function("!importBusy && drafts.length === 2", timeout=60000)
        ok(app.page.locator('[data-ocr]').count() >= 1 and app.js("drafts[0].total") is None, "ページ番号だけの文字があるPDFにOCRのボタンが出ない")
        app.page.click('[data-ocr="0"]')
        app.page.wait_for_function("drafts[0].ocr === true", timeout=240000)
        d2 = app.js("({c: [drafts[0].staff, drafts[0].family, drafts[0].other, drafts[0].total], w: drafts[0].warnings.join(' / ')})")
        eq(d2["c"][:3], [12, 6, 2], "ページ番号つきの画像PDFをOCRで読んだ受診者数")
        ok(d2["c"][3] == 20 or "合計が一致しません" in d2["w"], f"計を読み違えたときは知らせる: {d2}")
        # 設定の「OCRの動作確認」
        app.tab("settings")
        app.page.click("#ocrTest")
        app.page.wait_for_function("/✅|⚠|❌/.test(document.querySelector('#ocrTestSt').textContent)", timeout=120000)
        ok("✅" in app.page.inner_text("#ocrTestSt"), "OCRの動作確認: " + app.page.inner_text("#ocrTestSt"))
        eq(app.external, [], "OCRで外部に通信した")
    finally:
        app.close()


@test("AI分析：分析項目を複数選んで、1つの依頼文にまとめられる")
def t_ai_multi(app):
    app.js("refreshMonths('2026-08')")
    app.tab("ai")
    cards = app.js("[...document.querySelectorAll('#aiPrompts [data-pk]')].map(b => b.dataset.pk).filter(k => k !== 'none')")
    ok(len(cards) >= 3, "分析項目のボタン")
    first = app.js("document.querySelector('#aiPrompts .ai-pcard.on').dataset.pk")
    other = next(k for k in cards if k != first)
    app.page.click(f'#aiPrompts [data-pk="{other}"]')
    eq(sorted(app.js("[...document.querySelectorAll('#aiPrompts .ai-pcard.on')].map(b => b.dataset.pk)")), sorted([first, other]), "2つ選ばれている")
    pt = app.page.input_value("#aiPromptText")
    ok("観点1" in pt and "観点2" in pt and "全体のまとめ" in pt, "複数の観点の依頼文")
    ok("観点1" in app.page.input_value("#aiText"), "AIに渡す文章に依頼が入っていない")
    ok("・" in app.page.input_value("#aiAnsTitle"), "回答の題名に2つの項目名")
    # もう一度押すと外れる
    app.page.click(f'#aiPrompts [data-pk="{other}"]')
    eq(app.js("[...document.querySelectorAll('#aiPrompts .ai-pcard.on')].map(b => b.dataset.pk)"), [first], "押し直すと外れる")
    ok("観点1" not in app.page.input_value("#aiPromptText"), "1つに戻ったのに観点の見出しが残る")


@test("取り込み：読み取りを一時停止・再開できる")
def t_pause(browser, base):
    app = App(browser, base)
    try:
        import shutil
        src = fixtures.build(tempfile.mkdtemp(prefix="geppou_pause_src_"))
        folder = tempfile.mkdtemp(prefix="geppou_pause_")
        pdfs = [f for f in sorted(os.listdir(src)) if f.endswith(".pdf")]
        files = []
        for k in range(40):
            f = pdfs[k % len(pdfs)]
            dst = os.path.join(folder, f"{k:02d}_{f}")
            shutil.copy(os.path.join(src, f), dst)
            files.append(dst)
        app.tab("import")
        app.page.check('input[name="mmode"][value="auto"]')
        app.page.set_input_files("#files", files)
        app.page.click("#start")
        app.page.wait_for_function("importBusy")
        app.page.click("#pauseRead")
        ok("再開" in app.page.inner_text("#pauseRead"), "一時停止のあとに「再開」にならない")
        app.page.wait_for_timeout(1500)
        n1 = app.js("Number((document.querySelector('#progress').textContent.match(/(\\d+) \\/ 40/) || [0, -1])[1])")
        app.page.wait_for_timeout(1500)
        n2 = app.js("Number((document.querySelector('#progress').textContent.match(/(\\d+) \\/ 40/) || [0, -1])[1])")
        ok(n1 == n2 and n1 < 40 and app.js("importBusy"), f"一時停止中に読み進んでいる（{n1}→{n2}）")
        ok("一時停止中" in app.page.inner_text("#progress"), "一時停止中の表示")
        app.page.click("#pauseRead")
        app.page.wait_for_function("!importBusy", timeout=120000)
        eq(app.js("drafts.length"), 40, "再開後に全部読み取る")
        ok(app.page.locator("#pauseRead").is_hidden(), "読み終えたのにボタンが残る")
    finally:
        app.close()


@test("AI分析：推移は上の対象月ではなく最新の月まで。保存した回答に期間が入る")
def t_ai_month(app):
    app.js("refreshMonths('2026-07')")
    app.tab("ai")
    app.page.click('[data-aiper="0"]')
    ok("2026年7月" in app.page.inner_text("#aiHeadM"), "今月のみは上の対象月")
    app.page.click('[data-aiper="12"]')
    ok(app.page.inner_text("#aiHeadM").endswith("2026年8月の推移"), "推移は最新の月まで: " + app.page.inner_text("#aiHeadM"))
    app.page.fill("#aiAns", "テストの回答です。")
    app.page.fill("#aiAnsTitle", "期間のテスト")
    app.page.click("#aiAnsSave")
    app.page.wait_for_function("(state.aiNotes || []).some(x => x.title === '期間のテスト') && /期間のテスト/.test(document.querySelector('#aiNotes').textContent)")
    eq(app.js("(() => { const x = state.aiNotes.find(x => x.title === '期間のテスト'); return [x.month, x.periodLabel]; })()"), ["2026-08", "2026年7月〜2026年8月の推移"], "保存した回答の月と期間")
    ok("2026年7月〜2026年8月の推移" in app.page.inner_text("#aiNotes"), f"保存した回答の一覧に期間: {app.page.inner_text('#aiNotes')[:200]} {app.errors}")
    app.js("state.aiNotes = state.aiNotes.filter(x => x.title !== '期間のテスト'); saveState(); refreshMonths('2026-08')")
    app.page.click('[data-aiper="0"]')


@test("サマリーのAI：分割しない／3分割／5分割を選べ、分けたときは窓を並べて窓ごとにコピーできる")
def t_ais_split(browser, base):
    app = App(browser, base, "?demo=1")
    try:
        app.tab("summary")
        app.page.click("#aisStart")
        app.page.click('[data-aissplit="1"]')
        posts = app.js("reportsOf(curMonth()).map(r => r.post)")
        ok(len(posts) > 50, "デモの公館数")
        eq(app.js("document.querySelectorAll('#aisParts textarea').length"), 1, "分割しないときの窓の数")
        whole = app.page.input_value("#aisText")
        ok(f"【{posts[0]}】" in whole and f"【{posts[-1]}】" in whole, "全文に最初と最後の公館が入っていない")
        for n in (3, 5):
            app.page.click(f'[data-aissplit="{n}"]')
            texts = app.js("[...document.querySelectorAll('#aisParts textarea')].map(t => t.value)")
            eq(len(texts), n, f"{n}分割の窓の数")
            ok(all(t.startswith("【取扱注意】") for t in texts), "どの窓にも注意書き")
            ok(all(sum(f"【{p}】" in t for t in texts) == 1 for p in posts), f"{n}分割で、どの公館もどれか1つの窓だけに入る")
            eq(app.js("document.querySelectorAll('[data-ais-go=\"gaixer\"]').length"), n, "窓ごとのコピーのボタン")
        # 確認のチェックで、全部の窓のボタンが押せるようになる。窓ごとにコピーできる
        app.page.check("#aisOk")
        ok(app.js("[...document.querySelectorAll('[data-ais-go]')].every(b => !b.disabled)"), "確認後にすべての窓のボタンが押せない")
        app.js("window.open = () => null; window._copied = []; copyText = async (t) => { _copied.push(t); return true; }; 0")
        app.page.click('[data-ais-go="gaixer"][data-pi="2"]')
        app.page.wait_for_function("_copied.length === 1")
        ok("✓ 3つ目" in app.page.inner_text('[data-ais-go="gaixer"][data-pi="2"]'), "3つ目のボタンの表示")
        ok(app.js("_copied[0]") == app.js("document.querySelector('#aisText2').value"), "3つ目の窓の文章がコピーされていない")
        # Excelで書き出す：依頼文のシートと、1行1公館の公館データのシート（5つの窓の全部）
        with app.page.expect_download() as dl:
            app.page.click("#aisXlsx")
        import openpyxl
        ok(dl.value.suggested_filename.endswith(".xlsx"), "Excelのファイル名")
        xp = os.path.join(tempfile.mkdtemp(prefix="geppou_xlsx_"), "summary.xlsx")
        dl.value.save_as(xp)
        wb = openpyxl.load_workbook(xp)
        eq(wb.sheetnames, ["依頼文", "公館データ"], "Excelのシート")
        guide = "\n".join(str(r[0].value or "") for r in wb["依頼文"].iter_rows())
        ok("管理番号：AI-" in guide and "取扱注意" in guide and "公館ごとに" in guide, "依頼文のシートの中身")
        data = [[c.value for c in r] for r in wb["公館データ"].iter_rows()]
        eq(data[0], ["組", "公館", "地域", "AIに渡す内容"], "公館データの見出し")
        eq(sorted(r[1] for r in data[1:]), sorted(posts), "公館データに全公館（1行1公館）")
        ok(all("受診者数" in (r[3] or "") for r in data[1:]), "各公館の内容")
        app.page.click('[data-aissplit="1"]')
        eq(app.js("document.querySelectorAll('#aisParts textarea').length"), 1, "分割しないに戻す")
    finally:
        app.close()


@test("取り込み：壊れたファイル・空のファイルは、日本語の理由を付けて失敗にする（止まらない）")
def t_bad_files(browser, base):
    app = App(browser, base)
    try:
        folder = tempfile.mkdtemp(prefix="geppou_bad_")
        open(os.path.join(folder, "空_202607.pdf"), "wb").close()
        open(os.path.join(folder, "壊れた_202607.pdf"), "wb").write(b"%PDF-1.4 broken")
        open(os.path.join(folder, "偽_202607.pdf"), "wb").write(b"PK\x03\x04garbage")
        app.tab("import")
        app.page.check('input[name="mmode"][value="auto"]')
        app.page.set_input_files("#files", [os.path.join(folder, f) for f in sorted(os.listdir(folder))])
        app.page.click("#start")
        app.page.wait_for_function("!importBusy && drafts.length === 3", timeout=60000)
        w = app.js("drafts.map(d => d.warnings.join(' '))")
        ok(all("読み込み失敗" in x for x in w), f"失敗の印: {w}")
        ok(not any(re.search(r"[A-Za-z]{6,}", x.replace("PDF", "")) for x in w), f"英語のエラーが残っている: {w}")
        ok(any("空です" in x for x in w), "空のファイルの理由")
        eq(app.errors, [], "画面のエラー")
    finally:
        app.close()


@test("不具合報告：宛先・件名・本文（画面・版・エラーの記録）の下書きを作り、月報の内容・氏名を入れない")
def t_bug_report(app):
    ok(app.page.locator("#bugBtn").count() == 1, "不具合報告のボタンがない")
    body = app.js("bugReportText()")
    for w in ["・画面（タブ）：", "・したこと：", "版：", "ブラウザ：", "エラーの記録"]:
        ok(w in body, f"本文に「{w}」がない")
    names = {o for _, o, *_ in fixtures.REPORTS.values()} | {n for _, n, _ in fixtures.ROSTER if n}
    ok(not any(nm in body or nm.replace(" ", "") in body for nm in names), "本文に氏名が入っている")
    ok(app.js("BUG_TO") == "yusuke.hoshino@mofa.go.jp", "宛先")


@test("使い方（チュートリアル）：同期→フォルダ→ドラッグ＆ドロップ→保存の順に説明し、画面の場所を光らせる")
def t_tutorial(app):
    app.page.click("#tutBtn")
    ok(app.js("document.querySelector('#tut').open"), "使い方が開かない")
    titles = []
    for k in range(6):
        titles.append(app.page.inner_text("#tutBody h3"))
        app.page.click("#tutNext")
    ok("同期" in titles[0] and "エクスプローラ" in titles[1] and "ドラッグ" in titles[3] and "保存" in titles[4], f"手順の順番: {titles}")
    ok(not app.js("document.querySelector('#tut').open"), "おわりで閉じない")
    body = app.js("TUT_STEPS.map(s => s.html).join(' ')")
    ok("imukanmd" not in body and "iims" not in body.lower() and "http" not in body, "公開のページに組織内の情報（パスワード・サーバー名・アドレス）が入っている")
    app.page.click("#tutBtn"); app.page.click('[data-tut="3"]'); app.page.click("#tutShow")
    ok(app.js("page") == "import" and app.js("document.querySelector('#drop').classList.contains('tut-hl')"), "ドラッグの場所が光らない")
    app.page.click("#tutBack")
    ok(app.js("document.querySelector('#tut').open") and "4 / 6" in app.page.inner_text("#tutStepNo"), "使い方に戻れない")
    app.page.click("#tutClose")


@test("AI分析：AIに渡す文章をExcel（依頼と本文・表ごとのシート）で保存できる")
def t_ai_xlsx(app):
    app.js("refreshMonths('2026-08')")
    app.tab("ai")
    app.page.click('[data-aiper="12"]')
    ok(app.page.is_disabled("#aiXlsx"), "確認前にExcelで保存できる")
    app.page.check("#aiOk")
    with app.page.expect_download() as dl:
        app.page.click("#aiXlsx")
    ok(dl.value.suggested_filename.endswith(".xlsx"), "ファイル名")
    import openpyxl
    xp = os.path.join(tempfile.mkdtemp(prefix="geppou_aix_"), "ai.xlsx")
    dl.value.save_as(xp)
    wb = openpyxl.load_workbook(xp)
    ok(wb.sheetnames[0] == "依頼と本文" and len(wb.sheetnames) >= 3, f"シート: {wb.sheetnames}")
    body = "\n".join(str(r[0].value or "") for r in wb["依頼と本文"].iter_rows())
    ok("管理番号：AI-" in body and "取扱注意" in body, "依頼と本文の中身")
    sh = wb[wb.sheetnames[1]]
    ok(sh.max_row >= 2 and sh.max_column >= 2, "表のシートに行と列がある")
    names = {o for _, o, *_ in fixtures.REPORTS.values()}
    allv = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
    ok(not any(nm in allv or nm.replace(" ", "") in allv for nm in names), "Excelに氏名が入っている")


@test("国名の一覧は50音順（漢字の国名も読みで並ぶ）")
def t_gojuon(app):
    names = app.js("ALL_PLACES")
    for a, b in [("ウルグアイ", "英国"), ("英国", "エクアドル"), ("カメルーン", "韓国"), ("韓国", "ガンビア"), ("ケニア", "広州総領事館"), ("中央アフリカ", "中国"),
                 ("ボリビア", "香港総領事館"), ("香港総領事館", "ホンジュラス"), ("ミクロネシア", "南アフリカ共和国"), ("パングラデシュ".replace("パ", "バ"), "東ティモール")]:
        ok(names.index(a) < names.index(b), f"{a} が {b} より前にない")
    opts = app.js("(() => { const d = document.createElement('div'); d.innerHTML = postDatalist(); return [...d.querySelectorAll('option')].map(o => o.value); })()")
    eq(opts, [x for x in names if x != "日本"], "公館の候補も50音順の1つの一覧")


@test("AI分析：対象の地域は北米と中南米を分ける")
def t_ai_regions(app):
    app.js("refreshMonths('2026-08')")
    app.tab("ai")
    regs = app.js("[...document.querySelectorAll('[data-aireg]')].map(b => b.textContent)")
    ok("北米" in regs and "中南米" in regs and not any("米州" in r for r in regs), f"地域のボタン: {regs}")
    ok("アジア" in regs and "大洋州" in regs and "アジア・大洋州" not in regs, f"アジアと大洋州: {regs}")
    eq(app.js("['オーストラリア', 'フィジー', 'パプアニューギニア', 'フィリピン', 'インド', '東ティモール'].map(aiRegionOf)"), ["大洋州", "大洋州", "大洋州", "アジア", "アジア", "アジア"], "アジアと大洋州の公館")
    eq(app.js("normScope({regions: ['アジア・大洋州'], posts: []}).regions"), ["アジア", "大洋州"], "以前の「アジア・大洋州」の選択")
    # 説明資料のボタン（財務省・人事課）：押すと依頼文が入る
    for k, w in [("mof", "財務省"), ("jinji", "人事課")]:
        app.page.click(f'[data-pk="{k}"]')
        ok(w in app.page.input_value("#aiText") and "数値は集計にあるものだけ" in app.page.input_value("#aiText"), f"{w}向けの依頼文が入らない")
    # 地域のボタンで、その地域の国が下の欄でまとめて選ばれる／外れる。「選択中」の札は出さない
    on = "[...document.querySelectorAll('[data-aipost].on')].map(b => b.dataset.aipost).sort()"
    app.page.click('[data-aireg="アフリカ"]')
    eq(app.js(on), ["ウガンダ", "ケニア"], "アフリカを押すとアフリカの国が選ばれる")
    ok(app.js("document.querySelector('[data-aireg=\"アフリカ\"]').classList.contains('on')"), "地域のボタンが選択の表示にならない")
    eq(app.js("document.querySelectorAll('[data-unreg], [data-unpost], #aiSel').length"), 0, "「選択中」の表示が残っている")
    app.page.click('[data-aipost="ケニア"]')
    eq(app.js(on), ["ウガンダ"], "国を個別に外す")
    ok(app.js("document.querySelector('[data-aireg=\"アフリカ\"]').classList.contains('part')"), "一部だけ選ばれている地域の表示")
    ok("ウガンダ" in app.page.input_value("#aiText") and "ケニア" not in app.page.input_value("#aiText").split("## ")[0], "選んだ国で集計")
    app.page.click('[data-aireg="アフリカ"]')
    eq(app.js(on), ["ウガンダ", "ケニア"], "一部のときに地域を押すと全部選ぶ")
    app.page.click('[data-aireg="アフリカ"]')
    eq(app.js(on), [], "もう一度押すとまとめて外れる")
    # 全世界：すべての国が選ばれ、もう一度押すとすべて外れる
    app.page.click('[data-aireg=""]')
    eq(app.js(on), ["ウガンダ", "ケニア", "タイ", "フィリピン"], "全世界ですべての国を選ぶ")
    ok(app.js("document.querySelector('[data-aireg=\"\"]').classList.contains('on')") and "全世界" in app.page.inner_text("#aiPostSel"), "全世界の表示")
    ok("・全世界・" in app.page.input_value("#aiText"), "全世界で集計")
    app.page.click('[data-aireg=""]')
    eq(app.js(on), [], "全世界をもう一度押すとすべて外れる")
    eq(app.js("['ホンジュラス', 'ペルー', 'メキシコ'].map(aiRegionOf)"), ["中南米", "中南米", "中南米"], "中南米の公館")
    eq(app.js("D.posts.filter(p => p.iso3 === 'USA' || p.iso3 === 'CAN').every(p => aiRegionOf(p.name) === '北米')"), True, "北米の公館")
    eq(app.js("normScope({regions: ['米州'], posts: []}).regions"), ["北米", "中南米"], "以前の「米州」の選択")
    t = app.js("(() => { state.reports['2026-08|ペルー'] = {month: '2026-08', post: 'ペルー', region: '米州', total: 7, staff: 4, family: 2, other: 1, diseases: [], hei: ''}; "
               "const t = aiText('2026-08', new Set(['visits']), {regions: ['中南米'], posts: []}); delete state.reports['2026-08|ペルー']; return t; })()")
    ok("中南米" in t and "7" in t and "ケニア" not in t, f"中南米だけの集計: {t[:200]}")


@test("PDFレポート：パスワード imukanmd で暗号化され、開ける")
def t_pdf(app):
    from pypdf import PdfReader
    b64 = app.js("""async () => {
        const blob = await buildReportPdf('2026-08', new Set(REPORT_SECTIONS.map(s => s[0])));
        const u8 = new Uint8Array(await blob.arrayBuffer()); let s = '';
        for (let i = 0; i < u8.length; i += 0x8000) s += String.fromCharCode(...u8.subarray(i, i + 0x8000));
        return btoa(s); }""")
    r = PdfReader(io.BytesIO(base64.b64decode(b64)))
    ok(r.is_encrypted, "暗号化されていない")
    ok(r.decrypt("wrong-password") == 0, "違うパスワードで開けてしまう")
    ok(r.decrypt(fixtures.PASSWORD) != 0, "imukanmd で開けない")
    ok(len(r.pages) >= 3, f"ページ数が少ない: {len(r.pages)}")


@test("PDFレポート：緊急移送・邦人対応は報告のあるものだけ載せる")
def t_pdf_cases(app):
    # サマリー（15行には「緊急移送：特記なし」の行がある）を除いた章で確かめる
    text = "async (m) => (await reportBlocks(m, new Set(REPORT_SECTIONS.map(s => s[0]).filter(k => k !== 'summary')))).map(e => e.textContent).join('\\n')"
    aug = app.js(text, "2026-08")
    for w in ["緊急移送", "邦人対応", "報告なし"]:
        ok(w not in aug, f"報告のない8月に「{w}」が載っている")
    aug_all = app.js("async (m) => (await reportBlocks(m, new Set(REPORT_SECTIONS.map(s => s[0])))).map(e => e.textContent).join('\\n')", "2026-08")
    ok("7. 公館別 5行サマリー" in aug_all and "6. 巡回健診" in aug_all, "章の番号が詰まっていない（統計分析が5、緊急移送なしで巡回健診が6）")
    jul = app.js(text, "2026-07")
    ok("緊急移送症例" in jul and "邦人対応事例" in jul and "南アフリカ共和国" in jul, "7月の緊急移送・邦人対応が載っていない")
    only = app.js("""async () => { const keep = state.cases; state.cases = {'2026-07': casesOf('2026-07').filter(c => c.type === 'hogo')};
        const t = (await reportBlocks('2026-07', new Set(['cases']))).map(e => e.textContent).join('\\n'); state.cases = keep; return t; }""")
    ok("1. 邦人対応" in only and "緊急移送症例" not in only and "緊急移送・" not in only, "邦人対応だけのとき緊急移送の見出しが残る")


@test("PDFレポート：期間（3か月など）と地域を選んで作れる")
def t_pdf_period(app):
    from pypdf import PdfReader
    all_secs = "new Set(REPORT_SECTIONS.map(s => s[0]))"
    txt = lambda span, regs: app.js(f"""async () => (await withRegions({regs}, () => {span} !== 1 ? periodBlocks('2026-08', {span}, {all_secs}, {{label: reportLabel('2026-08', {span}), scope: regionsLabel({regs})}})
        : reportBlocks('2026-08', {all_secs}, {{label: reportLabel('2026-08'), scope: regionsLabel({regs})}}))).map(e => e.textContent).join('\\n')""")
    t3 = txt(3, "[]")
    ok("2026年6月〜2026年8月（3か月）" in t3, "3か月の表紙")
    ok("延べ81名" in t3, "3か月の受診者数の合計（7月35名＋8月46名）")
    ok("ケニア" in t3 and "フィリピン" in t3 and "提出月数" in t3, "公館別の期間の表")
    ok("緊急移送症例" in t3 and "南アフリカ共和国" in t3, "期間中の緊急移送（7月）")
    ta = txt(3, "['アフリカ']")
    ok("対象：アフリカ" in ta and "延べ56名" in ta, "アフリカだけの合計（ケニア・ウガンダ）")
    ok("フィリピン" not in ta and "タイ" not in ta, "アフリカ以外の公館が入っている")
    t1 = txt(1, "['アジア']")
    ok("2026年8月分" in t1 and "対象：アジア" in t1 and "フィリピン" in t1 and "ケニア" not in t1, "1か月・地域のレポート")
    # 個別の国を選ぶ（地域と組み合わせもできる）
    tp = app.js("""async () => (await withRegions([], () => periodBlocks('2026-08', 3, new Set(REPORT_SECTIONS.map(s => s[0])), {label: reportLabel('2026-08', 3), scope: regionsLabel([], ['ケニア'])}), ['ケニア'])).map(e => e.textContent).join('\\n')""")
    ok("対象：ケニア" in tp and "延べ38名" in tp and "ウガンダ" not in tp and "フィリピン" not in tp, "ケニアだけの3か月（18名＋20名）")
    tm = app.js("""async () => (await withRegions(['アジア'], () => reportBlocks('2026-08', new Set(['visits']), {label: reportLabel('2026-08'), scope: regionsLabel(['アジア'], ['ウガンダ'])}), ['ウガンダ'])).map(e => e.textContent).join('\\n')""")
    ok("フィリピン" in tm and "ウガンダ" in tm and "ケニア" not in tm and "アジア・ウガンダ" in tm, "地域と国の組み合わせ")
    eq(app.js("reportFileName('2026-08', 3, ['大洋州'], ['ケニア', 'ウガンダ', 'タイ'])"), "医務官月報統合レポート_202608_3か月_大洋州・ケニア・ウガンダほか1.pdf", "国を選んだときのファイル名")
    t0 = txt(0, "[]")
    # 地域で絞っても、全期間はすべての月報の最初の月から（フィリピンは8月だけでも、7月からの2か月）
    ok("全期間（2026年7月〜2026年8月・2か月）" in txt(0, "['アジア']"), "地域で絞ったときの全期間")
    ok("全期間（2026年7月〜2026年8月・2か月）" in t0, f"全期間の表紙: {t0[:120]}")
    ok("延べ81名（月平均41名）であった。" in t0, "全期間の受診者数（前の期間とは比べない）")
    eq(app.js("reportFileName('2026-08', 0, [])"), "医務官月報統合レポート_202608_全期間.pdf", "全期間のファイル名")
    eq(app.js("Object.keys(state.reports).length"), 6, "作成後に元のデータに戻っている")
    b64 = app.js("""async () => {
        const blob = await buildReportPdf('2026-08', new Set(REPORT_SECTIONS.map(s => s[0])), null, {span: 6, regions: ['アフリカ']});
        const u8 = new Uint8Array(await blob.arrayBuffer()); let s = '';
        for (let i = 0; i < u8.length; i += 0x8000) s += String.fromCharCode(...u8.subarray(i, i + 0x8000));
        return btoa(s); }""")
    r = PdfReader(io.BytesIO(base64.b64decode(b64)))
    ok(r.decrypt(fixtures.PASSWORD) != 0 and len(r.pages) >= 3, "6か月・アフリカのPDF")
    eq(app.js("reportFileName('2026-08', 6, ['北米', '中南米'])"), "医務官月報統合レポート_202608_6か月_北米・中南米.pdf", "ファイル名")
    # 画面：期間と地域のボタン
    app.tab("report")
    app.page.click('[data-rspan="0"]')
    ok("全期間" in app.page.inner_text("#rspanLab"), "全期間の表示")
    # 公館別サマリーの行数（はじめは5行）
    ok("on" in app.page.get_attribute('[data-rsum="5"]', "class"), "PDFのサマリーの行数のはじめが5行でない")
    app.page.click('[data-rsum="10"]')
    eq(app.js("loadPrefs().reportSumLines"), 10, "選んだ行数を覚える")
    app.page.click('[data-rsum="5"]')
    app.page.click('[data-rspan="36"]')
    ok("36か月" in app.page.inner_text("#rspanLab"), "3年の表示: " + app.page.inner_text("#rspanLab"))
    app.page.click('[data-rspan="12"]')
    app.page.click('[data-rreg="中南米"]')
    ok("12か月" in app.page.inner_text("#rspanLab") and "中南米" in app.page.inner_text("#rspanLab"), "選んだ期間・地域の表示")
    app.page.click("#rpost")
    eq(app.js("document.querySelectorAll('.combo-list .combo-item').length"), 4, "レポートの国の欄の候補（月報のある4か国）")
    app.page.keyboard.press("ArrowDown"); app.page.keyboard.press("Enter")
    ok(app.page.inner_text("#rpostChips").strip() != "", "▼とEnterで国を追加できない")
    app.page.click('[data-rreg=""]')
    app.page.fill("#rpost", "Philippines")
    app.page.click("#rpostAdd")
    ok("フィリピン" in app.page.inner_text("#rpostChips") and "フィリピン" in app.page.inner_text("#rspanLab"), "国を追加（英語名から）")
    app.page.click('[data-rspan="1"]'); app.page.click('[data-rreg=""]')
    eq(app.page.inner_text("#rpostChips"), "", "全世界で国の選択も外れる")
    # 「直近」は上で選んでいる月ではなく、月報のある最新の月まで
    app.js("refreshMonths('2026-07')"); app.tab("report")
    app.page.click('[data-rspan="3"]')
    ok("直近 2026年6月〜2026年8月（3か月）" in app.page.inner_text("#rspanLab"), "直近3か月は最新の月まで: " + app.page.inner_text("#rspanLab"))
    # 期間を指定（数字だけで入れられる。前後が逆でも入れ替える）
    app.page.click('[data-rspan="custom"]')
    ok(app.page.is_visible("#rfrom"), "期間を指定の欄が出ない")
    app.page.fill("#rfrom", "202608"); app.page.press("#rfrom", "Tab")
    app.page.fill("#rto", "202607"); app.page.press("#rto", "Tab")
    ok("2026年7月〜2026年8月（2か月）" in app.page.inner_text("#rspanLab"), "指定した期間: " + app.page.inner_text("#rspanLab"))
    app.page.click('[data-rspan="1"]')
    ok(not app.page.is_visible("#rfrom") and "2026年8月分" in app.page.inner_text("#rspanLab"), "直近1か月に戻す")
    app.js("refreshMonths('2026-08')")


@test("PDFレポート：統計分析（選んだ統計）とAI分析の結果を入れられる")
def t_pdf_stats_ai(app):
    # AI分析タブで回答を保存
    app.js("refreshMonths('2026-08')")
    app.tab("ai")
    app.page.fill("#aiAnsTitle", "今月の全体像（テスト）")
    app.page.fill("#aiAns", "※この集計は外務省AI（MOFA GaiXer）・政府AI源内専用です。それ以外のサービスに入力した場合は、直ちに履歴を削除してください。\n## 結論\n- 受診者数は**46名**であった。\n| 地域 | 件数 |\n|---|---|\n| アフリカ | 31 |\n今後も注意が必要である。")
    ok(not app.page.is_disabled("#aiAnsSave"), "保存ボタンが押せない")
    app.page.click("#aiAnsSave")
    app.page.wait_for_function("(state.aiNotes || []).length === 1")
    nid = app.js("state.aiNotes[0].id")
    t = app.js(f"""async () => (await reportBlocks('2026-08', new Set(['stats', 'ai']), {{label: reportLabel('2026-08'), scope: '全世界', stats: ['desc', 'share', 'change'], aiNotes: ['{nid}']}})).map(e => e.textContent).join('\\n')""")
    for w in ["統計分析", "記述統計（1公館1か月あたりの受診者数）", "構成比（地域別）", "前月と比べて増えた・減った公館", "AI分析の結果", "今月の全体像（テスト）", "受診者数は46名であった。", "アフリカ"]:
        ok(w in t, f"レポートに「{w}」がない")
    for w in ["ばらつき（地域別）", "相関", "専用です。それ以外"]:
        ok(w not in t, f"選んでいない・載せない「{w}」が入っている")
    t2 = app.js("async () => (await reportBlocks('2026-08', new Set(['ai']), {label: reportLabel('2026-08'), scope: '全世界', aiNotes: []})).map(e => e.textContent).join('\\n')")
    ok("AI分析の結果" not in t2, "回答を選んでいないのにAI分析の章がある")
    # 画面：統計の選択とAIの回答の選択
    app.tab("report")
    if not app.page.is_checked('[data-rsec="stats"]'): app.page.check('[data-rsec="stats"]')
    if not app.page.is_checked('[data-rsec="ai"]'): app.page.check('[data-rsec="ai"]')
    ok(app.page.locator("[data-rstat]").count() == 9 and app.page.is_visible("#rstatPick"), "統計のチェックが出ない")
    ok(app.page.is_checked(f'[data-rai="{nid}"]'), "その月のAIの回答が最初から選ばれていない")
    app.page.uncheck('[data-rsec="stats"]')
    ok(not app.page.is_visible("#rstatPick"), "統計分析を外しても統計のチェックが出ている")
    app.page.check('[data-rsec="stats"]')
    app.js("state.aiNotes = []")


@test("バックアップ：暗号化して書き出し、元どおりに戻せる")
def t_backup(app):
    same, enc = app.js("""async () => {
        const txt = await encryptBackup(state);
        const back = await decryptBackup(txt);
        return [JSON.stringify(back) === JSON.stringify(state), txt]; }""")
    ok(same, "復号した内容が元と違う")
    ok("架空" not in enc and "ケニア" not in enc, "バックアップの中身が平文で読める")
    msg = app.js("decryptBackup(JSON.stringify({format:'geppou-backup-encrypted', salt:'AAAA', iv:'AAAA', data:'AAAA'})).then(() => '', e => e.message)")
    ok("復号できません" in msg, "壊れたバックアップでエラーにならない")


# ----------------------------------------------------------------------- デモ・通信

@test("デモ表示：全タブがエラーなく表示でき、実データに書き込まない")
def t_demo(browser, base):
    app = App(browser, base, "?demo=1")
    try:
        app.page.wait_for_function("demoMode && Object.keys(state.reports).length > 50", timeout=30000)
        for name in app.js("[...document.querySelectorAll('#tabs button[data-page]')].map(b => b.dataset.page)"):
            before = len(app.errors)
            app.tab(name)
            eq(app.errors[before:], [], f"デモ「{name}」タブのエラー")
        eq(app.js("idb.get('state').then(s => s ? Object.keys(s.reports).length : 0)"), 0, "デモのデータが保存領域に書かれている")
        eq(app.external, [], "デモ表示中の外部通信")
    finally:
        app.close()


@test("地図：国・首都・都市のカードを開いても外部に接続しない（Wikipedia の自動読み込み・旅行リンクなし）")
def t_place_card(app):
    app.tab("map")
    app.page.wait_for_function("!!document.getElementById('wmap') && !!document.getElementById('wmap').data")
    before = len(app.external)
    html = app.js("""(() => { const c = PLACES.capitals.find(r => r[0] === 'ナイロビ') || PLACES.capitals[0];
        openPlaceCard({kind: 'capital', name: c[0], iso: c[4], pop: c[3], en: c[5]});
        openPlaceCard({kind: 'country', name: 'ブータン', iso: 'BTN'});
        return document.querySelector('#postcard').innerHTML; })()""")
    app.page.wait_for_timeout(800)
    eq(app.external[before:], [], "カードを開いたときの外部通信")
    for word in ["google.com/travel", "booking.com", "expedia", "wikisum"]:
        ok(word not in html, f"カードに「{word}」が残っている")
    ok("ja.wikipedia.org/wiki/" in html, "Wikipedia へのリンク（押したときだけ開く）がない")
    app.js("openPostCard(['ケニア'])")
    html = app.js("document.querySelector('#postcard').innerHTML")
    ok("google.com/travel" not in html and "booking.com" not in html, "公館カードに旅行リンクが残っている")
    eq(app.js("document.querySelectorAll('canvas.sattex, .mapctl [data-skin], #tabs [data-page=whm]').length"), 0, "衛星写真・医療情報マップが残っている")


@test("外部への通信：取り込みから出力まで一度もない")
def t_network(app):
    eq(app.external, [], "外部への通信")


# ----------------------------------------------------------------------- 実行

def main():
    only = sys.argv[1] if len(sys.argv) > 1 else ""
    srv = start_server()
    base = f"http://127.0.0.1:{srv.server_address[1]}"
    folder = fixtures.build(tempfile.mkdtemp(prefix="geppou_test_"))
    failed = 0
    t0 = time.time()
    with sync_playwright() as p:
        try:
            browser = p.chromium.launch()
        except Exception:
            browser = p.chromium.launch(executable_path="/opt/pw-browsers/chromium")
        app = App(browser, base)
        try:
            app.drafts = import_fixtures(app, folder)
        except Exception:
            print("✗ 見本の月報を取り込めませんでした")
            traceback.print_exc()
            return 1
        for name, fn in TESTS:
            if only and only not in name:
                continue
            try:
                if fn.__code__.co_argcount == 2:
                    fn(browser, base)
                else:
                    fn(app)
                print(f"✓ {name}")
            except Exception as e:
                failed += 1
                print(f"✗ {name}\n    {e if isinstance(e, AssertionError) else traceback.format_exc()}")
        if app.errors:
            failed += 1
            print("✗ 画面のエラー:\n    " + "\n    ".join(app.errors))
        app.close()
        browser.close()
    srv.shutdown()
    n = sum(1 for name, _ in TESTS if only in name)
    print(f"\n{n - failed if failed <= n else 0}/{n} 件 成功（{time.time() - t0:.0f}秒）" + ("" if not failed else f"　失敗 {failed}件"))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
