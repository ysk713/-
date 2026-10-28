"""自動テスト用の見本ファイル（月報・配置表）を作る。氏名・数値はすべて架空。

答えが分かっている内容にしてあり、テストではアプリの読み取り・集計結果がこの答えと一致するかを確かめる。
"""
import io
import os
import sys

from pypdf import PdfReader, PdfWriter

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "sample"))
from make_sample import kou_text, make_real_pdf, make_zip  # noqa: E402

PASSWORD = "imukanmd"

# 公館の表記, 医務官, (職員, 家族, その他, 計), 乙表の病名と件数, 丙表
REPORTS = {
    "ケニア_202607_月報.pdf": ("在ケニア日本国大使館", "架空 一郎", "2026年7月", (10, 5, 3, 18),
        [("急性鼻咽頭炎［かぜ］", 4), ("胃腸炎", 3), ("デング熱", 2), ("健康診断", 5)],
        "1. 医療事情\n市内でコレラ症例の報告が続いている。飲料水の注意喚起を館員に周知した。\n"
        "2. 邦人保護\n邦人旅行者1名が交通外傷で入院し、南アフリカの病院へ緊急移送した。手術を受け回復した。"),
    "ケニア_202608_月報.pdf": ("在ケニア日本国大使館", "見本 花子", "2026年8月", (12, 6, 2, 20),
        [("急性鼻咽頭炎［かぜ］", 6), ("湿疹", 2)],
        "1. 着任\n8月1日に新任医務官が着任した。\n2. 医療事情\nコレラの新規報告は減少傾向にある。"),
    "ウガンダ_202607_月報.pdf": ("在ウガンダ日本国大使館", "試験 次郎", "2026年7月", (5, 2, 0, 7),
        [("マラリア", 2), ("背部痛", 1)], "1. 医療事情\n特記事項なし。"),
    "2026年8月_ウガンダ.docx": ("在ウガンダ日本国大使館", "架空 一郎", "2026年8月", (7, 3, 1, 11),
        [("マラリア", 3), ("不眠症", 1)],
        "1. 着任\n8月5日、前任地ケニアより着任した。\n2. 巡回健診\n8月18日～21日 南スーダンへ巡回健診を実施し、館員12名を診察した。"),
    "フィリピン_202608_月報.pdf": ("在フィリピン日本国大使館", "検証 六子", "2026年8月", (9, 4, 2, 15),
        [("デング熱", 5), ("急性咽頭炎", 2)],
        "1. 感染症\n雨季入りに伴いデング熱が急増し、マニラ首都圏で病床がひっ迫している。"),
    # 合計欄が内訳の和と合わない（10 ≠ 3+2+1）→ 警告が出ること
    "タイ_202607_月報.pdf": ("在タイ日本国大使館", "試作 七美", "2026年7月", (3, 2, 1, 10),
        [("便秘", 1)], "1. 医療事情\nPM2.5は改善した。"),
}
# 形式: zip=ZIP形式のPDF（OCRテキスト入り）, real=普通のPDF, enc=パスワード付きPDF, docx=Word
FORMATS = {"ケニア_202607_月報.pdf": "zip", "ケニア_202608_月報.pdf": "zip", "ウガンダ_202607_月報.pdf": "real",
           "2026年8月_ウガンダ.docx": "docx", "フィリピン_202608_月報.pdf": "enc", "タイ_202607_月報.pdf": "zip"}

ROSTER_NAME = "医務官配置表_202608.xlsx"
# 8月の配置表。タイは配置ありなのに8月の月報がない → 未提出。スーダンは退避中（配置なし扱い）
ROSTER = [("在ケニア大使館", "見本 花子", ""), ("在ウガンダ大使館", "架空 一郎", ""), ("在フィリピン大使館", "検証 六子", ""),
          ("在タイ大使館", "試作 七美", ""), ("在スーダン大使館", "", "退避中")]


def otsu_text(emb, period, diseases):
    lines = ["疾病分類件数表", emb, f"（{period}分）"]
    for i, (name, c) in enumerate(diseases, 1):
        lines.append(f"{i} Ａ００－Ｂ９９ 感染症及び寄生虫症 {c} {name}({c})")
    lines.append("報告者　医務官")
    return "\n".join(lines)


def pages_of(fname):
    emb, officer, period, (s, f, o, t), dis, hei = REPORTS[fname]
    y, m = period.replace("月", "").split("年")
    kou = kou_text(emb, officer, f"{y}-{int(m):02d}", (s, f, o))
    kou = kou.replace(f"受診者数　{s}　{f}　{o}　{s + f + o}", f"受診者数　{s}　{f}　{o}　{t}")
    return [kou, otsu_text(emb, period, dis), f"報告事項等\n{hei}"]


def make_docx(pages, path):
    from docx import Document
    from docx.enum.text import WD_BREAK
    doc = Document()
    for i, t in enumerate(pages):
        for line in t.split("\n"):
            doc.add_paragraph(line)
        if i < len(pages) - 1:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.save(path)


def encrypt_pdf(path, password):
    r = PdfReader(path)
    w = PdfWriter()
    for p in r.pages:
        w.add_page(p)
    w.encrypt(user_password=password, owner_password=None, algorithm="AES-128")
    buf = io.BytesIO()
    w.write(buf)
    with open(path, "wb") as f:
        f.write(buf.getvalue())


def make_roster(path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "医務官配置表"
    ws.append(["医務官配置表（テスト用・架空）"])
    ws.append(["No", "地域", "駐在公館名", "氏名", "ふりがな", "備考"])
    for i, (post, name, note) in enumerate(ROSTER, 1):
        ws.append([i, "", post, name, "", note])
    wb.save(path)


def build(out):
    os.makedirs(out, exist_ok=True)
    for fname, fmt in FORMATS.items():
        path = os.path.join(out, fname)
        pages = pages_of(fname)
        if fmt == "zip":
            make_zip(pages, path)
        elif fmt == "docx":
            make_docx(pages, path)
        else:
            make_real_pdf(pages, path)
            if fmt == "enc":
                encrypt_pdf(path, PASSWORD)
    make_roster(os.path.join(out, ROSTER_NAME))
    return out


def roster_pdf_lines():
    # 表をPDFにしたときのテキスト（列は空白区切り。列の並びは Excel 版と違ってもよい）
    lines = ["医務官配置表（2026年8月1日現在・テスト用・架空）", "No 駐在公館名 氏名 ふりがな 備考"]
    kana = {"見本 花子": "みほん はなこ", "架空 一郎": "かくう いちろう", "検証 六子": "けんしょう むつこ", "試作 七美": "しさく ななみ"}
    for i, (post, name, note) in enumerate(ROSTER, 1):
        lines.append(f"{i} {post} {name} {kana.get(name, '')} {note}".strip())
    return lines


def make_roster_pdf(path):
    make_real_pdf(["\n".join(roster_pdf_lines())], path)


def make_image_pdf(pages, path):
    # 文字の情報がなく、画像だけのPDF（スキャンした月報の代わり）。OCR のテストに使う
    from PIL import Image, ImageDraw, ImageFont
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    font = ImageFont.truetype("/usr/share/fonts/opentype/ipafont-gothic/ipag.ttf", 34)
    c = canvas.Canvas(path, pagesize=A4)
    for k, t in enumerate(pages):
        img = Image.new("L", (1654, 2339), 255)
        d = ImageDraw.Draw(img)
        y = 120
        for line in t.split("\n"):
            d.text((110, y), line, font=font, fill=0)
            y += 56
        p = f"{path}.{k}.png"
        img.save(p)
        c.drawImage(p, 0, 0, *A4)
        c.showPage()
    c.save()
