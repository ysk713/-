"""統合Wordレポートの生成（python-docx + matplotlib）。"""
import io
import json
import math
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager as fm
from matplotlib.patches import Polygon
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import analysis, db, icd, maps, posts

DOC_FONT = "游ゴシック"

# ---------------------------------------------------------------- matplotlib 日本語フォント

_FONT_CANDIDATES = [
    "C:/Windows/Fonts/YuGothM.ttc", "C:/Windows/Fonts/meiryo.ttc", "C:/Windows/Fonts/msgothic.ttc",
    "/System/Library/Fonts/ヒラギノ角ゴシック W3.ttc", "/System/Library/Fonts/Hiragino Sans GB.ttc",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc", "/usr/share/fonts/noto-cjk/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/opentype/ipafont-gothic/ipag.ttf", "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
]


def setup_font():
    for path in _FONT_CANDIDATES:
        if os.path.exists(path):
            fm.fontManager.addfont(path)
            plt.rcParams["font.family"] = fm.FontProperties(fname=path).get_name()
            break
    plt.rcParams["axes.unicode_minus"] = False


setup_font()

REGION_COLORS = {"アフリカ": "#E69F00", "アジア・大洋州": "#D55E00", "ヨーロッパ": "#0072B2",
                 "中東": "#CC79A7", "米州": "#009E73", "未分類": "#999999"}


def _png(fig):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------- 図


def chart_region_visits(rep):
    g = rep.groupby("region")[["staff", "family", "other"]].sum().reindex(posts.REGIONS).fillna(0)
    fig, ax = plt.subplots(figsize=(9, 4.2))
    left = [0] * len(g)
    for col, label, c in [("staff", "職員", "#4472C4"), ("family", "家族", "#ED7D31"), ("other", "その他", "#A5A5A5")]:
        ax.barh(g.index, g[col], left=left, label=label, color=c)
        left = [a + b for a, b in zip(left, g[col])]
    for i, v in enumerate(left):
        ax.text(v, i, f" {int(v)}", va="center", fontsize=10)
    ax.invert_yaxis()
    ax.set_xlabel("受診者数（延べ）")
    ax.legend(loc="lower right")
    ax.spines[["top", "right"]].set_visible(False)
    ax.set_title("地域別 受診者数", fontsize=13)
    return _png(fig)


def chart_region_category(dis):
    d = dis[dis["category"] != icd.Z_CATEGORY]
    pv = d.pivot_table(index="region", columns="category", values="count", aggfunc="sum").fillna(0)
    pv = pv.reindex([r for r in posts.REGIONS if r in pv.index])
    order = pv.sum().sort_values(ascending=False).index
    fig, ax = plt.subplots(figsize=(10, 4.8))
    left = [0] * len(pv)
    for cat in order:
        ax.barh(pv.index, pv[cat], left=left, color=icd.COLORS.get(cat), label=icd.SHORT.get(cat, cat))
        left = [a + b for a, b in zip(left, pv[cat])]
    ax.invert_yaxis()
    ax.set_xlabel("診断件数（検診・予防接種等を除く）")
    ax.legend(ncol=4, fontsize=8, loc="upper center", bbox_to_anchor=(0.5, -0.15), frameon=False)
    ax.spines[["top", "right"]].set_visible(False)
    ax.set_title("地域別 疾病分類", fontsize=13)
    return _png(fig)


def chart_trend():
    r = db.reports()
    if r["month"].nunique() < 2:
        return None
    pv = r.pivot_table(index="month", columns="region", values="total", aggfunc="sum").fillna(0)
    fig, ax = plt.subplots(figsize=(9, 4))
    for reg in [x for x in posts.REGIONS if x in pv.columns]:
        ax.plot(pv.index, pv[reg], marker="o", label=reg, color=REGION_COLORS[reg])
    ax.plot(pv.index, pv.sum(axis=1), marker="s", color="black", lw=2, label="合計")
    ax.set_ylabel("受診者数")
    ax.legend(ncol=3, fontsize=9)
    ax.grid(alpha=0.3)
    ax.spines[["top", "right"]].set_visible(False)
    ax.set_title("受診者数の月次推移（地域別）", fontsize=13)
    return _png(fig)


def chart_world(rep, dis, ev, jk, tr):
    """matplotlib 版の世界地図（Word埋め込み用の静止画）。"""
    gj = maps.geojson()
    cd = maps.country_dominant(dis).set_index("iso3") if not dis.empty else None
    fig, ax = plt.subplots(figsize=(16, 8.4))
    used = {}
    for f in gj["features"]:
        if f["id"] == "ATA":
            continue
        col = "#ECEFF1"
        if cd is not None and f["id"] in cd.index:
            row = cd.loc[f["id"]]
            col = icd.COLORS.get(row["category"], "#ccc")
            if row["total"] < maps.MIN_TOTAL_FOR_FULL_COLOR:
                col = maps.blend_with_white(col)
            else:
                used[row["category"]] = icd.COLORS.get(row["category"])
        geom = f["geometry"]
        polys = geom["coordinates"] if geom["type"] == "MultiPolygon" else [geom["coordinates"]]
        for poly in polys:
            ax.add_patch(Polygon(poly[0], closed=True, fc=col, ec="white", lw=0.4))
    for j in jk.itertuples() if not jk.empty else []:
        a, b = maps.coord(j.from_post), maps.coord(j.dest)
        if a and b:
            ax.annotate("", xy=(b[1], b[0]), xytext=(a[1], a[0]),
                        arrowprops=dict(arrowstyle="-|>", color="#00897B", lw=1.8, connectionstyle="arc3,rad=0.25"))
    for t in tr:
        a, b = maps.coord(t.get("from_post")), maps.coord(t["post"])
        if a and b:
            ax.annotate("", xy=(b[1], b[0]), xytext=(a[1], a[0]),
                        arrowprops=dict(arrowstyle="-|>", color="#6A1B9A", lw=1.8, ls="--",
                                        connectionstyle="arc3,rad=-0.25"))
    if not rep.empty:
        r = rep.dropna(subset=["total"])
        pts = [(maps.coord(p), t) for p, t in zip(r["post"], r["total"]) if maps.coord(p)]
        ax.scatter([c[1] for c, _ in pts], [c[0] for c, _ in pts], s=[6 * t for _, t in pts],
                   facecolor="none", edgecolor="#263238", lw=1, alpha=0.7, zorder=5)
    if not ev.empty:
        mk = {"感染症": ("P", "#B00020"), "気象・災害": ("s", "#E67E22"), "治安・社会": ("^", "#5D4037"),
              "邦人保護": ("D", "#1565C0")}
        for k, (m, c) in mk.items():
            e = ev[ev["kind"] == k].drop_duplicates("post")
            pts = [maps.coord(p) for p in e["post"] if maps.coord(p)]
            if pts:
                off = list(mk).index(k)
                ax.scatter([p[1] + 2.4 * math.cos(off * 1.57) for p in pts],
                           [p[0] + 1.8 * math.sin(off * 1.57) for p in pts],
                           marker=m, s=60, color=c, edgecolor="white", lw=0.6, zorder=6, label=k)
    from matplotlib.patches import Patch
    from matplotlib.lines import Line2D
    h1 = [Patch(color=c, label=icd.SHORT.get(k, k)) for k, c in used.items()]
    h2 = [Line2D([], [], color="#00897B", lw=2, label="巡回健診"),
          Line2D([], [], color="#6A1B9A", lw=2, ls="--", label="医務官の異動"),
          Line2D([], [], marker="o", ls="", mfc="none", mec="#263238", label="受診者数(○の大きさ)")]
    leg = ax.legend(handles=h1, loc="lower left", fontsize=9, title="最多疾病分類（検診・予防接種除く）",
                    title_fontsize=9, framealpha=0.9)
    ax.add_artist(leg)
    hs, ls = ax.get_legend_handles_labels()
    ax.legend(handles=h2 + hs, loc="lower right", fontsize=9, framealpha=0.9)
    ax.set_xlim(-170, 190)
    ax.set_ylim(-58, 85)
    ax.set_aspect("equal")
    ax.set_axis_off()
    ax.set_title("世界地図で見る今月の状況", fontsize=16)
    return _png(fig)


# ---------------------------------------------------------------- docx ヘルパー


def _font(run, size=None, bold=None, color=None):
    run.font.name = DOC_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), DOC_FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _p(doc, text="", size=10.5, bold=False, align=None, color=None, style=None):
    p = doc.add_paragraph(style=style)
    if text:
        _font(p.add_run(text), size, bold, color)
    if align:
        p.alignment = align
    return p


def _h(doc, text, level=1):
    h = doc.add_heading(level=level)
    _font(h.add_run(text), {1: 16, 2: 13, 3: 11.5}[level], True, "1F3864")
    return h


def _table(doc, header, rows, widths=None, bold_rows=()):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        _font(c.paragraphs[0].add_run(str(h)), 9, True)
        _shade(c, "D9E2F3")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            _font(cells[i].paragraphs[0].add_run("" if v is None else str(v)), 9, ri in bold_rows)
            if ri in bold_rows:
                _shade(cells[i], "F2F2F2")
    if widths:
        t.autofit = False
        for i, w in enumerate(widths):
            t.columns[i].width = Cm(w)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return t


def _shade(cell, hex_fill):
    from docx.oxml import OxmlElement
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _n(v):
    try:
        return "―" if v is None or (isinstance(v, float) and math.isnan(v)) else f"{int(v)}"
    except (TypeError, ValueError):
        return "―"


# ---------------------------------------------------------------- 全体概況（自動）


def overview_bullets(month, rep, dis, ev, jk, tr, miss):
    out = []
    prv = db.reports(analysis.prev_month(month))
    tot = int(rep["total"].fillna(0).sum())
    s = f"提出{len(rep)}公館の受診者数は延べ{tot}名"
    if not prv.empty:
        s += f"（前月{int(prv['total'].fillna(0).sum())}名、提出{len(prv)}公館）"
    out.append(s + "。")
    if not dis.empty:
        d = dis[dis["category"] != icd.Z_CATEGORY].groupby("category")["count"].sum().sort_values(ascending=False)
        if len(d):
            out.append("疾病分類（検診・予防接種を除く）は " +
                       "、".join(f"{icd.SHORT.get(c, c)} {n}件" for c, n in d.head(3).items()) + " の順に多い。")
    for kind in ["感染症", "気象・災害", "治安・社会"]:
        e = ev[ev["kind"] == kind]
        if e.empty:
            continue
        g = e.groupby("keyword")["post"].apply(lambda x: "・".join(sorted(set(x))))
        out.append(f"{kind}：" + "／".join(f"{k}（{v}）" for k, v in g.items()) + "。")
    if not jk.empty:
        out.append(f"巡回健診の報告 {len(jk)}件：" +
                   "、".join(f"{a}→{b or '（行先不明）'}" for a, b in zip(jk["from_post"], jk["dest"])) + "。")
    trm = [t for t in tr if t["month"] == month]
    if trm:
        out.append("医務官の交代：" + "、".join(f"{t['post']}（{t['before']}→{t['after']}）" for t in trm) + "。")
    if not miss.empty:
        out.append(f"月報未提出 {len(miss)}公館：" + "、".join(miss["post"]) + "。")
    return out


# ---------------------------------------------------------------- 本体


def build(month, include_map=True):
    rep = db.reports(month)
    dis = db.diseases(month)
    ev = db.events(month)
    jk = db.junkai(month)
    tr = analysis.transfers()
    miss = analysis.missing_posts(month)
    y, m = month.split("-")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2))
    sec.top_margin = sec.bottom_margin = Cm(2)
    st = doc.styles["Normal"]
    st.font.name = DOC_FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), DOC_FONT)

    # 表紙
    for _ in range(6):
        _p(doc)
    _p(doc, "医務官月報 統合レポート", 26, True, WD_ALIGN_PARAGRAPH.CENTER, "1F3864")
    _p(doc, f"{y}年{int(m)}月分", 18, False, WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc)
    _p(doc, f"提出 {len(rep)}公館　／　未提出 {len(miss)}公館", 13, False, WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    sections = ["1. 全体概況", "2. 世界地図", "3. 受診者数統計（甲表）", "4. 疾病分類統計（乙表）",
                "5. 巡回健診・医務官の異動", "6. 公館別 5行サマリー"]
    if not include_map:
        sections.remove("2. 世界地図")
    _h(doc, "目次", 1)
    for s in sections:
        _p(doc, s, 11)

    _h(doc, "1. 全体概況", 1)
    for b in overview_bullets(month, rep, dis, ev, jk, tr, miss):
        _p(doc, b, 10.5, style="List Bullet")

    if include_map:
        _h(doc, "2. 世界地図", 1)
        doc.add_picture(chart_world(rep, dis, ev, jk, [t for t in tr if t["month"] == month]), width=Cm(17))
        _p(doc, "※国の塗り分けは最多疾病分類（実患8件未満の国は薄色）。○は受診者数、記号は丙表で報告された感染症・災害等。",
           8.5, color="555555")

    _h(doc, "3. 受診者数統計（甲表）", 1)
    if not rep.empty:
        doc.add_picture(chart_region_visits(rep), width=Cm(16))
    tc = chart_trend()
    if tc:
        doc.add_picture(tc, width=Cm(16))
    mm = analysis.mom(month).set_index("post")
    rows, bold = [], []
    allrows = list(rep.itertuples())
    for reg in posts.REGIONS + ["未分類"]:
        rr = [r for r in allrows if r.region == reg]
        mr = miss[miss["region"] == reg] if not miss.empty else miss
        if not rr and mr.empty:
            continue
        for r in rr:
            d = mm["diff"].get(r.post)
            rows.append([reg, r.post, r.officer or "", _n(r.staff), _n(r.family), _n(r.other), _n(r.total),
                         "" if d is None or d != d else f"{int(d):+d}"])
        for r in mr.itertuples():
            rows.append([reg, r.post, r.officer or "", "―", "―", "―", "―（未提出）", ""])
        bold.append(len(rows))
        rows.append(["", f"{reg} 小計", "", _n(sum(r.staff or 0 for r in rr)), _n(sum(r.family or 0 for r in rr)),
                     _n(sum(r.other or 0 for r in rr)), _n(sum(r.total or 0 for r in rr)), ""])
    bold.append(len(rows))
    rows.append(["", "総計", "", _n(rep["staff"].sum()), _n(rep["family"].sum()), _n(rep["other"].sum()),
                 _n(rep["total"].sum()), ""])
    _table(doc, ["地域", "公館", "医務官", "職員", "家族", "その他", "計", "前月比"], rows,
           [2.6, 3.4, 2.8, 1.3, 1.3, 1.3, 2.0, 1.5], bold_rows=bold)

    _h(doc, "4. 疾病分類統計（乙表）", 1)
    if not dis.empty:
        doc.add_picture(chart_region_category(dis), width=Cm(16.5))
        pv = dis.pivot_table(index="category", columns="region", values="count", aggfunc="sum").fillna(0)
        regs = [r for r in posts.REGIONS if r in pv.columns]
        pv = pv.reindex(columns=regs).assign(計=lambda x: x.sum(axis=1)).sort_values("計", ascending=False)
        _h(doc, "分類別・地域別件数", 2)
        _table(doc, ["疾病分類"] + regs + ["計"],
               [[icd.SHORT.get(c, c)] + [_n(v) for v in row] for c, row in pv.iterrows()])
        _h(doc, "診断名トップ10（検診・予防接種等を除く）", 2)
        top = dis[dis["category"] != icd.Z_CATEGORY].groupby("name").agg(
            件数=("count", "sum"), 公館数=("post", "nunique")).sort_values("件数", ascending=False).head(10)
        _table(doc, ["順位", "診断名", "分類", "件数", "公館数"],
               [[i + 1, n, icd.SHORT.get(icd.classify(n), ""), r.件数, r.公館数]
                for i, (n, r) in enumerate(top.iterrows())], [1.2, 6.5, 3.5, 1.8, 1.8])
    else:
        _p(doc, "乙表データなし")

    _h(doc, "5. 巡回健診・医務官の異動", 1)
    _h(doc, "巡回健診", 2)
    if jk.empty:
        _p(doc, "今月の巡回健診の報告なし")
    else:
        _table(doc, ["実施公館（出発）", "巡回先", "期間", "記載"],
               [[j.from_post, j.dest or "（不明）", j.period or "", j.note.rsplit(" [", 1)[0][:80]]
                for j in jk.itertuples()], [3.2, 3.2, 2.4, 8.2])
    _h(doc, "医務官の異動・交代", 2)
    trm = [t for t in tr if t["month"] == month]
    if not trm:
        _p(doc, "前月からの医務官の交代なし（月報ヘッダー・配置表による）")
    else:
        _table(doc, ["公館", "前任", "後任", "後任の前任地"],
               [[t["post"], t["before"], t["after"], t["from_post"] or ""] for t in trm])

    _h(doc, "6. 公館別 5行サマリー", 1)
    for reg in posts.REGIONS + ["未分類"]:
        rr = rep[rep["region"] == reg]
        mr = miss[miss["region"] == reg] if not miss.empty else miss
        if rr.empty and mr.empty:
            continue
        _h(doc, reg, 2)
        for r in rr.itertuples():
            _h(doc, f"{r.post}（医務官：{r.officer or '―'}）", 3)
            summ = r.summary or analysis.rule_summary(month, r.post)
            for line in summ.splitlines()[:5]:
                _p(doc, line, 10, style="List Bullet")
        for r in mr.itertuples():
            _h(doc, f"{r.post}（医務官：{r.officer or '―'}）", 3)
            _p(doc, f"{int(m)}月分の月報提出なし", 10, color="C00000")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
